"""C-26 evidence tests: APA-08 (abstract) and APA-09 (keywords) formatting.

Covers the style-first path (template styles present), the direct-formatting
fallback (styles absent — the defect that left the Abstract heading as
plain left-aligned Normal text), the italic-label rule, and the
Practitioner Notes removal that C-26 shipped alongside.
"""

import shutil
import zipfile

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from lxml import etree
from unittest.mock import patch

from app.domain import reporting_ownership as ro
from app.domain.canonical_jultp_template import CANONICAL_STRUCTURE
from app.services import output_generation_samfix as S
from app.services.abstract_keywords_formatting import (
    apply_abstract_keywords_formatting,
)
from app.services.apa7_formatting_checks import (
    check_abstract_format,
    check_keywords_format,
)
from app.services.apa7_validator import validate as apa_validate
from app.services.document_analysis_services import (
    extract_abstract,
    extract_keywords,
    front_page_summary,
    load_paragraphs,
)
from app.services.docx_format_helpers import (
    normalise_heading_text,
    paragraph_text,
    para_style_id,
)

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WQ = f"{{{W}}}"

# The real style definitions from the regenerated APA base template.
FM_STYLES_XML = f"""<w:styles xmlns:w="{W}">
  <w:style w:type="paragraph" w:styleId="HeadingFrontPage">
    <w:name w:val="Heading Front Page"/>
    <w:pPr><w:jc w:val="center"/></w:pPr>
    <w:rPr><w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/>
           <w:b/><w:sz w:val="22"/></w:rPr>
  </w:style>
  <w:style w:type="paragraph" w:styleId="FrontPageText">
    <w:name w:val="Front Page Text"/>
    <w:pPr><w:ind w:left="0" w:firstLine="0"/></w:pPr>
    <w:rPr><w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/>
           <w:sz w:val="22"/></w:rPr>
  </w:style>
  <w:style w:type="paragraph" w:styleId="Keywords">
    <w:name w:val="Keywords"/>
    <w:pPr><w:ind w:left="720" w:firstLine="0"/></w:pPr>
    <w:rPr><w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/>
           <w:sz w:val="22"/></w:rPr>
  </w:style>
  <w:style w:type="character" w:styleId="KeywordsLabel">
    <w:name w:val="Keywords Label"/>
    <w:rPr><w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/><w:i/></w:rPr>
  </w:style>
</w:styles>"""


def _add_front_matter_styles(src):
    """Merge the APA front-matter styles into the document's styles.xml."""
    out = str(src) + ".fmstyles.docx"
    with zipfile.ZipFile(str(src)) as z:
        members = {i.filename: z.read(i.filename) for i in z.infolist()}
    existing = etree.fromstring(members["word/styles.xml"])
    have = {
        s.find(f"{WQ}name").get(f"{WQ}val")
        for s in existing.findall(f"{WQ}style")
        if s.find(f"{WQ}name") is not None
    }
    for st in etree.fromstring(FM_STYLES_XML.encode()).findall(f"{WQ}style"):
        if st.find(f"{WQ}name").get(f"{WQ}val") not in have:
            existing.append(st)
    members["word/styles.xml"] = etree.tostring(
        existing, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for n, d in members.items():
            zout.writestr(n, d)
    return out


def _doc(tmp_path, abstract_para_count=1, inline_keywords=True,
         indent_abstract=True, name="in.docx"):
    """A front page with an Abstract block and a Keywords block."""
    doc = Document()
    doc.add_paragraph("A Study of Things")
    doc.add_paragraph("Abstract", style="Normal")
    for i in range(abstract_para_count):
        p = doc.add_paragraph(
            f"Abstract body sentence {i}." if abstract_para_count > 1
            else "This abstract summarises the study."
        )
        if indent_abstract:
            p.paragraph_format.first_line_indent = 720
    if inline_keywords:
        kw = doc.add_paragraph("Keywords: feedback, learning")
    else:
        doc.add_paragraph("Keywords", style="Normal")
        kw = doc.add_paragraph("feedback, learning")
    doc.add_paragraph("Introduction", style="Heading 1")
    doc.add_paragraph("Body text.")
    path = tmp_path / name
    doc.save(str(path))
    return path


def _paras(path):
    with zipfile.ZipFile(str(path)) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    return [el for el in root.find(f"{WQ}body") if el.tag == f"{WQ}p"]


def _find(paras, text):
    for p in paras:
        if normalise_heading_text(paragraph_text(p)) == text:
            return p
    return None


def _status(results, rule_id):
    for r in results:
        if r["rule_id"] == rule_id:
            return r["status"]
    return "(absent)"


# ── style-first path ───────────────────────────────────────────────────────

def test_abstract_heading_promoted_to_front_page_style(tmp_path):
    src = _add_front_matter_styles(_doc(tmp_path))
    out = tmp_path / "out.docx"

    _, actions = apply_abstract_keywords_formatting(src, str(out))

    assert any(a["kind"] == "abstract_heading_style" for a in actions)
    heading = _find(_paras(out), "abstract")
    assert para_style_id(heading) == "HeadingFrontPage"
    assert _status(check_abstract_format(str(out)), "APA008") == "pass"


def test_abstract_body_gets_front_page_text_style_and_no_indent(tmp_path):
    src = _add_front_matter_styles(_doc(tmp_path))
    out = tmp_path / "out.docx"

    _, actions = apply_abstract_keywords_formatting(src, str(out))

    kinds = {a["kind"] for a in actions}
    assert "abstract_body_style" in kinds
    assert "abstract_body_indent" in kinds
    paras = _paras(out)
    body = next(p for p in paras if "Abstract body sentence" in paragraph_text(p)
                or "This abstract summarises" in paragraph_text(p))
    assert para_style_id(body) == "FrontPageText"
    ind = body.find(f"{WQ}pPr/{WQ}ind")
    assert ind.get(f"{WQ}firstLine") == "0"


def test_inline_keywords_line_italicised_and_indented(tmp_path):
    src = _add_front_matter_styles(_doc(tmp_path))
    out = tmp_path / "out.docx"

    _, actions = apply_abstract_keywords_formatting(src, str(out))

    kinds = {a["kind"] for a in actions}
    assert "keywords_label_italic" in kinds
    assert "keywords_indent" in kinds
    assert _status(check_keywords_format(str(out)), "APA009") == "pass"


def test_only_the_label_is_italic_not_the_terms(tmp_path):
    src = _add_front_matter_styles(_doc(tmp_path))
    out = tmp_path / "out.docx"
    apply_abstract_keywords_formatting(src, str(out))

    kw = next(p for p in _paras(out)
              if paragraph_text(p).lower().startswith("keywords"))
    runs = [(r.find(f"{WQ}t").text or "",
             r.find(f"{WQ}rPr/{WQ}rStyle") is not None
             or r.find(f"{WQ}rPr/{WQ}i") is not None)
            for r in kw.iter(f"{WQ}r")
            if r.find(f"{WQ}t") is not None]
    assert runs, "keywords paragraph has no text runs"
    assert runs[0][1] is True, "the Keywords: label should be italic"
    assert all(italic is False for _, italic in runs[1:]), \
        f"keyword terms must not be italic: {runs}"


def test_heading_plus_terms_layout_italicises_heading_and_terms_line(tmp_path):
    src = _add_front_matter_styles(_doc(tmp_path, inline_keywords=False))
    out = tmp_path / "out.docx"

    _, actions = apply_abstract_keywords_formatting(src, str(out))

    kinds = {a["kind"] for a in actions}
    assert "keywords_style" in kinds
    assert "keywords_label_italic" in kinds
    assert _status(check_keywords_format(str(out)), "APA009") == "pass"


# ── direct-formatting fallback (styles absent) ────────────────────────────

def test_missing_styles_falls_back_to_tracked_direct_formatting(tmp_path):
    """The defect C-26 closes: with no 'Heading Front Page' style the pass
    must still make the heading bold + centred (tracked) rather than no-op."""
    src = _doc(tmp_path)          # no front-matter styles merged in
    out = tmp_path / "out.docx"

    _, actions = apply_abstract_keywords_formatting(src, str(out))

    kinds = {a["kind"] for a in actions}
    assert "abstract_heading_bold" in kinds
    assert "abstract_heading_centre" in kinds
    heading = _find(_paras(out), "abstract")
    assert heading.find(f"{WQ}r/{WQ}rPr/{WQ}b") is not None
    p_pr = heading.find(f"{WQ}pPr")
    assert p_pr.find(f"{WQ}jc").get(f"{WQ}val") == "center"
    assert p_pr.find(f"{WQ}pPrChange") is not None
    assert _status(check_abstract_format(str(out)), "APA008") == "pass"


def test_idempotent_second_run(tmp_path):
    src = _add_front_matter_styles(_doc(tmp_path))
    o1 = tmp_path / "o1.docx"
    apply_abstract_keywords_formatting(src, str(o1))
    o2 = tmp_path / "o2.docx"
    _, actions = apply_abstract_keywords_formatting(str(o1), str(o2))
    assert actions == []


def test_no_abstract_or_keywords_is_a_noop(tmp_path):
    doc = Document()
    doc.add_paragraph("Body only.")
    src = tmp_path / "in.docx"
    doc.save(str(src))
    out = tmp_path / "out.docx"
    _, actions = apply_abstract_keywords_formatting(str(src), str(out))
    assert actions == []
    assert out.exists(), "a no-op pass must still write output_path"


def test_pprchange_stays_last_when_indent_follows(tmp_path):
    """w:pPrChange must remain the last child of w:pPr (Word rejects otherwise),
    so a later silent property has to slot in front of it."""
    src = _doc(tmp_path)
    out = tmp_path / "out.docx"
    apply_abstract_keywords_formatting(src, str(out))
    for p in _paras(out):
        p_pr = p.find(f"{WQ}pPr")
        if p_pr is None or p_pr.find(f"{WQ}pPrChange") is None:
            continue
        assert etree.QName(list(p_pr)[-1]).localname == "pPrChange"


# ── keyword extraction ────────────────────────────────────────────────────

def test_inline_keywords_extracted_and_label_stripped(tmp_path):
    path = _doc(tmp_path, inline_keywords=True)
    assert extract_keywords(load_paragraphs(str(path))) == ["feedback", "learning"]


def test_heading_form_keywords_extracted(tmp_path):
    path = _doc(tmp_path, inline_keywords=False)
    assert extract_keywords(load_paragraphs(str(path))) == ["feedback", "learning"]


def test_abstract_body_not_swallowed_by_keywords(tmp_path):
    path = _doc(tmp_path)
    ab = extract_abstract(load_paragraphs(str(path)))
    assert ab["word_count"] == 5  # "This abstract summarises the study."


# ── Practitioner Notes removal (C-26 / APA-20) ────────────────────────────

def test_legacy_practitioner_block_terminates_abstract(tmp_path):
    """A leftover PN block must stop the abstract scan (it is a terminator,
    not a requirement) — otherwise its notes inflate the word count."""
    doc = Document()
    doc.add_paragraph("Abstract", style="Heading 1")
    doc.add_paragraph("Short abstract.", style="Normal")
    doc.add_paragraph("Practitioner Notes", style="Normal")
    for i in range(5):
        doc.add_paragraph(f"Note {i} about practice.", style="Normal")
    doc.add_paragraph("Keywords", style="Normal")
    path = tmp_path / "doc.docx"
    doc.save(str(path))

    ab = extract_abstract(load_paragraphs(str(path)))
    assert ab["paragraph_count"] == 1
    assert "practice" not in " ".join(p.text for p in ab["paragraphs"])


def test_no_practitioner_notes_keys_in_canonical_structure():
    front = CANONICAL_STRUCTURE["front_page"]
    assert "practitioner_notes_count" not in front
    assert "practitioner_notes_heading" not in front


def test_front_page_summary_has_no_practitioner_key(tmp_path):
    path = _doc(tmp_path)
    assert "practitioner_notes_heading_found" not in \
        front_page_summary(load_paragraphs(str(path)))


def test_practitioner_rules_absent_from_validator(tmp_path):
    ids = {r["rule_id"] for r in apa_validate(str(_doc(tmp_path)))["results"]}
    for gone in ("FP005", "FP009", "FP012"):
        assert gone not in ids


# ── reporting ownership ───────────────────────────────────────────────────

@pytest.mark.parametrize("rule_id", ["APA008", "APA009"])
def test_front_matter_rules_are_sam_fixed_and_listed(rule_id):
    assert rule_id in ro.SAM_FIXED_RULE_IDS
    entry = ro.FAMILY_VERDICTS["APA"]
    assert rule_id in (entry[2] if len(entry) > 2 else [])


def test_fp005_not_sam_fixed_anymore():
    assert "FP005" not in ro.SAM_FIXED_RULE_IDS


# ── end-to-end ────────────────────────────────────────────────────────────

def _run_build(src, out):
    with patch.object(S, "call_llm_json", side_effect=Exception("skip llm")):
        with patch.object(S, "build_body_edit_plan", return_value={"action": "none", "edits": []}):
            with patch.object(
                S, "build_editorial_review_comment_plan",
                return_value={"action": "none", "comments": []},
            ):
                S.build_edited_document(str(src), str(out))


def test_full_build_reaches_apa08_and_apa09(tmp_path):
    src = _add_front_matter_styles(_doc(tmp_path))
    before = apa_validate(str(src))["results"]
    assert _status(before, "APA008") == "fail"
    assert _status(before, "APA009") == "fail"

    out = tmp_path / "out.docx"
    _run_build(src, out)

    after = apa_validate(str(out))["results"]
    assert _status(after, "APA008") == "pass"
    assert _status(after, "APA009") == "pass"