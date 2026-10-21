import shutil
import zipfile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from lxml import etree

from app.services.output_generation_samfix import (
    _apply_apa_body_geometry,
    _apply_apa_title_page,
    _make_page_number_header,
    build_apa_title_page_plan,
)

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WQ = f"{{{W}}}"
NS = {"w": W}


def _make_student_doc(path):
    doc = Document()
    for name in ("Article Title", "Authors", "Author Affiliations"):
        doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph("a study of AI", style="Article Title")
    doc.add_paragraph("Jane Doe", style="Authors")
    doc.add_paragraph("University of Example", style="Author Affiliations")
    doc.add_paragraph("PSY 101: Introduction to Psychology")
    doc.add_paragraph("Dr. Smith")
    doc.add_paragraph("October 10, 2026")
    doc.add_paragraph("Introduction", style="Heading 1")
    doc.add_paragraph("Body text.")
    doc.save(path)


def _run_title_page_pass(source, output, mode=None):
    geometry = output.with_name("geometry.docx")
    assert _apply_apa_body_geometry(str(source), str(geometry), 9600) is True
    assert _apply_apa_title_page(str(geometry), str(output), mode=mode) is True


def _body_paragraphs(path):
    with zipfile.ZipFile(path) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    return root.xpath("//w:body/w:p", namespaces=NS)


def _paragraph_text(paragraph):
    return "".join(node.text or "" for node in paragraph.iter(f"{WQ}t"))


def test_student_title_page_is_centered_and_has_no_banner_or_textbox(tmp_path):
    source = tmp_path / "input.docx"
    output = tmp_path / "output.docx"
    _make_student_doc(source)

    _run_title_page_pass(source, output, mode="student")
    paragraphs = _body_paragraphs(output)
    title = next(p for p in paragraphs if _paragraph_text(p) == "A Study of AI")

    assert title.find(f"{WQ}pPr/{WQ}pStyle").get(f"{WQ}val") == "ArticleTitle"
    assert title.find(f"{WQ}pPr/{WQ}jc").get(f"{WQ}val") == "center"
    assert sum(_paragraph_text(p) == "" for p in paragraphs[: paragraphs.index(title)]) >= 3

    with zipfile.ZipFile(output) as archive:
        document_xml = archive.read("word/document.xml")
        assert b"txbxContent" not in document_xml
        assert all(
            b"blip" not in archive.read(name)
            for name in archive.namelist()
            if name.startswith("word/header") and name.endswith(".xml")
        )
        assert all(
            b"PAGE" in archive.read(name)
            for name in archive.namelist()
            if name.startswith("word/header") and name.endswith(".xml")
        )


def test_professional_mode_adds_running_head_and_removes_student_metadata(tmp_path):
    source = tmp_path / "input.docx"
    output = tmp_path / "output.docx"
    shutil.copy("app/domain/APA7 Base Template.docx", source)

    _apply_apa_title_page(str(source), str(output), mode="professional")
    paragraphs = _body_paragraphs(output)
    text = [_paragraph_text(p) for p in paragraphs]
    assert "PSY 101: Introduction to Psychology" not in text
    assert "Dr. Instructor Name" not in text
    assert "October 10, 2026" not in text

    with zipfile.ZipFile(output) as archive:
        headers = [
            archive.read(name)
            for name in archive.namelist()
            if name.startswith("word/header") and name.endswith(".xml")
        ]
        assert headers
        assert all(b"TITLE OF PAPER" in header for header in headers)
        assert all(b"PAGE" in header for header in headers)
        assert all(b"blip" not in header for header in headers)


def test_mode_override_beats_auto_detection(tmp_path):
    source = tmp_path / "input.docx"
    _make_student_doc(source)

    plan = build_apa_title_page_plan(str(source), mode="professional")
    assert plan["mode"] == "professional"

    plan = build_apa_title_page_plan(str(source), mode="student")
    assert plan["mode"] == "student"


def _header_with_running_head_and_logo():
    hdr = etree.Element(f"{WQ}hdr", nsmap={"w": W})

    running = etree.SubElement(hdr, f"{WQ}p")
    run = etree.SubElement(running, f"{WQ}r")
    text = etree.SubElement(run, f"{WQ}t")
    text.text = "Running Head: TITLE OF PAPER"

    logo = etree.SubElement(hdr, f"{WQ}p")
    logo_run = etree.SubElement(logo, f"{WQ}r")
    logo_text = etree.SubElement(logo_run, f"{WQ}t")
    logo_text.text = "LOGO TEXT"

    return etree.tostring(hdr, xml_declaration=True, encoding="UTF-8", standalone=True)


def _header_paragraphs(header_xml):
    return etree.fromstring(header_xml).findall(f"{WQ}p")


def test_make_page_number_header_strips_existing_header_content():
    result = _make_page_number_header(_header_with_running_head_and_logo())
    paragraphs = _header_paragraphs(result)

    assert len(paragraphs) == 1
    text = _paragraph_text(paragraphs[0])
    assert "Running Head" not in text
    assert "LOGO" not in text
    assert "PAGE" in "".join(
        f.get(f"{WQ}instr", "") or "" for f in paragraphs[0].iter(f"{WQ}fldSimple")
    )


def test_make_page_number_header_is_right_aligned():
    result = _make_page_number_header(_header_with_running_head_and_logo())
    paragraph = _header_paragraphs(result)[0]
    jc = paragraph.find(f"{WQ}pPr/{WQ}jc")
    assert jc is not None and jc.get(f"{WQ}val") == "right"


def test_make_page_number_header_strips_running_head_sharing_page_paragraph():
    hdr = etree.Element(f"{WQ}hdr", nsmap={"w": W})
    para = etree.SubElement(hdr, f"{WQ}p")
    p_pr = etree.SubElement(para, f"{WQ}pPr")
    jc = etree.SubElement(p_pr, f"{WQ}jc")
    jc.set(f"{WQ}val", "left")
    run = etree.SubElement(para, f"{WQ}r")
    text = etree.SubElement(run, f"{WQ}t")
    text.text = "Running Head: TITLE OF PAPER"
    field = etree.SubElement(para, f"{WQ}fldSimple")
    field.set(f"{WQ}instr", " PAGE ")

    header_xml = etree.tostring(hdr, xml_declaration=True, encoding="UTF-8", standalone=True)
    result = _make_page_number_header(header_xml)

    paragraphs = _header_paragraphs(result)
    assert len(paragraphs) == 1
    assert "Running Head" not in _paragraph_text(paragraphs[0])
    jc = paragraphs[0].find(f"{WQ}pPr/{WQ}jc")
    assert jc.get(f"{WQ}val") == "right"
    assert "PAGE" in "".join(
        f.get(f"{WQ}instr", "") or "" for f in paragraphs[0].iter(f"{WQ}fldSimple")
    )


def test_body_geometry_blanks_header_and_preserves_footer(tmp_path):
    source = tmp_path / "input.docx"
    output = tmp_path / "output.docx"

    doc = Document()
    for name in ("Article Title", "Authors", "Author Affiliations"):
        doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph("A Study of AI", style="Article Title")
    doc.add_paragraph("Jane Doe", style="Authors")
    doc.add_paragraph("University of Example", style="Author Affiliations")
    doc.add_paragraph("Introduction", style="Heading 1")
    doc.add_paragraph("Body text.")

    header = doc.sections[0].header
    header.is_linked_to_previous = False
    header.paragraphs[0].text = "Running Head: PRE-EXISTING HEADER"

    footer = doc.sections[0].footer
    footer.is_linked_to_previous = False
    footer.paragraphs[0].text = "FOOTER SURVIVES"

    doc.save(source)

    assert _apply_apa_body_geometry(str(source), str(output), 9600) is True

    with zipfile.ZipFile(output) as archive:
        header_names = [
            n for n in archive.namelist()
            if n.startswith("word/header") and n.endswith(".xml")
        ]
        footer_names = [
            n for n in archive.namelist()
            if n.startswith("word/footer") and n.endswith(".xml")
        ]

        assert header_names
        for name in header_names:
            paragraphs = _header_paragraphs(archive.read(name))
            assert len(paragraphs) == 1
            assert "PRE-EXISTING" not in _paragraph_text(paragraphs[0])
            assert "Running Head" not in _paragraph_text(paragraphs[0])

        assert footer_names
        for name in footer_names:
            assert b"FOOTER SURVIVES" in archive.read(name)
