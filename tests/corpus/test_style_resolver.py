"""Tests for tools/corpus/style_resolver.py (K-0.4a: basedOn chain + theme fonts,
K-0.4b: docDefaults, Word built-in defaults, numeric units).

Each test builds a small .docx at test time with hand-written styles.xml (and
theme fonts where needed), so the expected answer is known exactly.
Run:  python -m pytest tests/corpus/test_style_resolver.py -v
"""

import zipfile

import pytest
from docx import Document
from lxml import etree

from tools.corpus.docx_view import accepted_view
from tools.corpus.fixtures import build_good_docx
from tools.corpus.style_resolver import (
    UNSET, WORD_DEFAULT, get_resolver, resolve_effective, resolve_paragraph,
)

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _make_docx(path, body_xml, styles_xml, major=None, minor=None):
    """A real .docx (python-docx) with our own <w:body>, styles.xml and theme fonts."""
    Document().save(path)
    with zipfile.ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist()}

    parts["word/document.xml"] = (
        f'<w:document xmlns:w="{W_NS}"><w:body>{body_xml}</w:body></w:document>'
    ).encode()
    parts["word/styles.xml"] = f'<w:styles xmlns:w="{W_NS}">{styles_xml}</w:styles>'.encode()

    theme = etree.fromstring(parts["word/theme/theme1.xml"])
    for kind, family in (("majorFont", major), ("minorFont", minor)):
        if family:
            theme.find(f".//{{{A_NS}}}{kind}/{{{A_NS}}}latin").set("typeface", family)
    parts["word/theme/theme1.xml"] = etree.tostring(theme, xml_declaration=True,
                                                    encoding="UTF-8", standalone=True)

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in parts.items():
            z.writestr(name, data)
    return accepted_view(path)


def _style(sid, inner, type_="paragraph", default=False, based_on=None):
    attrs = f'w:type="{type_}" w:styleId="{sid}"' + (' w:default="1"' if default else "")
    based = f'<w:basedOn w:val="{based_on}"/>' if based_on else ""
    return f"<w:style {attrs}><w:name w:val=\"{sid}\"/>{based}{inner}</w:style>"


def _p(text="Text", ppr="", rpr=""):
    ppr = f"<w:pPr>{ppr}</w:pPr>" if ppr else ""
    rpr = f"<w:rPr>{rpr}</w:rPr>" if rpr else ""
    return f'<w:p>{ppr}<w:r>{rpr}<w:t xml:space="preserve">{text}</w:t></w:r></w:p>'


def _first_run(view):
    return view.paragraphs[0].runs[0]


# ---------------- backlog acceptance ----------------

def test_spacing_only_on_normal_resolves_480(tmp_path):
    styles = _style("Normal", '<w:pPr><w:spacing w:line="480" w:lineRule="auto"/></w:pPr>',
                    default=True)
    view = _make_docx(tmp_path / "a.docx", _p(), styles)
    eff = resolve_effective(_first_run(view), view)
    assert eff["spacing"]["line"] == "480"
    assert eff["spacing"]["lineRule"] == "auto"
    assert eff["sources"]["spacing.line"] == "style:Normal"


@pytest.mark.parametrize("theme_attr, expected", [
    ("minorHAnsi", "Georgia"),
    ("minorAscii", "Georgia"),
    ("majorHAnsi", "Arial"),
])
def test_theme_font_read_from_theme1(tmp_path, theme_attr, expected):
    styles = _style("Normal", f'<w:rPr><w:rFonts w:asciiTheme="{theme_attr}"/></w:rPr>',
                    default=True)
    view = _make_docx(tmp_path / "t.docx", _p(), styles, major="Arial", minor="Georgia")
    eff = resolve_effective(_first_run(view), view)
    assert eff["font"] == expected
    assert eff["sources"]["font"] == f"style:Normal (theme {theme_attr})"


# ---------------- chain behaviour ----------------

def test_basedon_chain_merges_attributes(tmp_path):
    styles = (
        _style("Normal", '<w:pPr><w:spacing w:line="480" w:lineRule="auto"/></w:pPr>'
                         '<w:rPr><w:sz w:val="22"/></w:rPr>', default=True)
        + _style("Heading1", '<w:pPr><w:spacing w:before="0"/><w:jc w:val="center"/></w:pPr>',
                 based_on="Normal")
    )
    view = _make_docx(tmp_path / "b.docx", _p("Method", '<w:pStyle w:val="Heading1"/>'), styles)
    eff = resolve_effective(_first_run(view), view)
    assert eff["spacing"]["line"] == "480"
    assert eff["sources"]["spacing.line"] == "style:Normal"
    assert eff["spacing"]["before"] == "0"
    assert eff["sources"]["spacing.before"] == "style:Heading1"
    assert eff["align"] == "center"
    assert eff["size"] == "22"


def test_direct_formatting_beats_style(tmp_path):
    styles = _style("Normal", '<w:pPr><w:jc w:val="left"/></w:pPr><w:rPr><w:sz w:val="22"/></w:rPr>',
                    default=True)
    body = _p(ppr='<w:jc w:val="right"/>', rpr='<w:sz w:val="28"/>')
    view = _make_docx(tmp_path / "d.docx", body, styles)
    eff = resolve_effective(_first_run(view), view)
    assert (eff["size"], eff["sources"]["size"]) == ("28", "direct")
    assert (eff["align"], eff["sources"]["align"]) == ("right", "direct")


def test_character_style_chain(tmp_path):
    styles = (
        _style("Normal", '<w:rPr><w:rFonts w:ascii="Calibri"/><w:sz w:val="22"/></w:rPr>',
               default=True)
        + _style("BaseChar", '<w:rPr><w:sz w:val="24"/></w:rPr>', type_="character")
        + _style("Emph", '<w:rPr><w:rFonts w:ascii="Arial"/></w:rPr>', type_="character",
                 based_on="BaseChar")
    )
    view = _make_docx(tmp_path / "c.docx", _p(rpr='<w:rStyle w:val="Emph"/>'), styles)
    eff = resolve_effective(_first_run(view), view)
    assert (eff["font"], eff["sources"]["font"]) == ("Arial", "style:Emph")
    assert (eff["size"], eff["sources"]["size"]) == ("24", "style:BaseChar")


def test_theme_attribute_beats_explicit_font_on_same_element(tmp_path):
    styles = _style("Normal", '<w:rPr><w:rFonts w:ascii="Arial" w:asciiTheme="minorHAnsi"/></w:rPr>',
                    default=True)
    view = _make_docx(tmp_path / "e.docx", _p(), styles, minor="Georgia")
    assert resolve_effective(_first_run(view), view)["font"] == "Georgia"


def test_hanging_and_firstline_are_one_group(tmp_path):
    styles = _style("Normal", '<w:pPr><w:ind w:firstLine="720"/></w:pPr>', default=True)
    body = _p(ppr='<w:ind w:left="720" w:hanging="720"/>')
    view = _make_docx(tmp_path / "h.docx", body, styles)
    indent = resolve_effective(_first_run(view), view)["indent"]
    assert indent == {"left": "720", "right": "0", "firstLine": UNSET, "hanging": "720"}


def test_start_beats_style_left(tmp_path):
    styles = _style("Normal", '<w:pPr><w:ind w:left="1440"/></w:pPr>', default=True)
    view = _make_docx(tmp_path / "s.docx", _p(ppr='<w:ind w:start="360"/>'), styles)
    assert resolve_effective(_first_run(view), view)["indent"]["left"] == "360"


def test_basedon_loop_does_not_hang(tmp_path):
    styles = (_style("A", '<w:pPr><w:jc w:val="center"/></w:pPr>', based_on="B")
              + _style("B", "", based_on="A"))
    view = _make_docx(tmp_path / "l.docx", _p(ppr='<w:pStyle w:val="A"/>'), styles)
    eff = resolve_effective(_first_run(view), view)
    assert eff["align"] == "center"
    assert [sid for sid, _ in get_resolver(view).chain("A")] == ["A", "B"]


def test_word_defaults_when_nothing_sets_a_value(tmp_path):
    view = _make_docx(tmp_path / "u.docx", _p(), _style("Normal", "", default=True))
    eff = resolve_effective(_first_run(view), view)
    assert (eff["font"], eff["size"], eff["align"]) == ("Times New Roman", "20", "left")
    assert eff["spacing"] == {"line": "240", "lineRule": "auto", "before": "0", "after": "0"}
    assert eff["indent"] == {"left": "0", "right": "0", "firstLine": "0", "hanging": "0"}
    assert set(eff["sources"].values()) == {WORD_DEFAULT}
    assert eff["units"]["size_pt"] == 10.0
    assert eff["units"]["line_multiple"] == 1.0


def test_paragraph_without_runs(tmp_path):
    styles = _style("Normal", '<w:pPr><w:spacing w:line="480"/></w:pPr>', default=True)
    view = _make_docx(tmp_path / "n.docx", "<w:p/>", styles)
    assert resolve_paragraph(view.paragraphs[0], view)["spacing"]["line"] == "480"


# ---------------- K-0.4b: docDefaults ----------------

def _doc_defaults(rpr="", ppr=""):
    return (f"<w:docDefaults><w:rPrDefault><w:rPr>{rpr}</w:rPr></w:rPrDefault>"
            f"<w:pPrDefault><w:pPr>{ppr}</w:pPr></w:pPrDefault></w:docDefaults>")


def test_docdefaults_used_when_no_style_sets_value(tmp_path):
    styles = (_doc_defaults(rpr='<w:sz w:val="24"/>',
                            ppr='<w:spacing w:after="200" w:line="276" w:lineRule="auto"/>')
              + _style("Normal", "", default=True))
    view = _make_docx(tmp_path / "dd.docx", _p(), styles)
    eff = resolve_effective(_first_run(view), view)
    assert (eff["size"], eff["sources"]["size"]) == ("24", "docDefaults")
    assert (eff["spacing"]["line"], eff["sources"]["spacing.line"]) == ("276", "docDefaults")
    assert eff["units"]["line_multiple"] == 1.15
    assert eff["units"]["after_pt"] == 10.0
    assert eff["sources"]["spacing.before"] == WORD_DEFAULT   # docDefaults didn't set it


def test_style_beats_docdefaults(tmp_path):
    styles = (_doc_defaults(ppr='<w:spacing w:line="276"/>')
              + _style("Normal", '<w:pPr><w:spacing w:line="480"/></w:pPr>', default=True))
    view = _make_docx(tmp_path / "sd.docx", _p(), styles)
    eff = resolve_effective(_first_run(view), view)
    assert (eff["spacing"]["line"], eff["sources"]["spacing.line"]) == ("480", "style:Normal")


def test_docdefaults_theme_font(tmp_path):
    styles = (_doc_defaults(rpr='<w:rFonts w:asciiTheme="minorHAnsi" w:hAnsiTheme="minorHAnsi"/>')
              + _style("Normal", "", default=True))
    view = _make_docx(tmp_path / "dt.docx", _p(), styles, minor="Georgia")
    eff = resolve_effective(_first_run(view), view)
    assert eff["font"] == "Georgia"
    assert eff["sources"]["font"] == "docDefaults (theme minorHAnsi)"


# ---------------- K-0.4b: numeric units ----------------

def test_units_auto_spacing_size_and_indent(tmp_path):
    styles = _style("Normal", '<w:pPr><w:spacing w:line="480" w:lineRule="auto" w:before="240"/>'
                              '<w:ind w:left="1440" w:firstLine="720"/><w:jc w:val="start"/></w:pPr>'
                              '<w:rPr><w:sz w:val="22"/></w:rPr>', default=True)
    view = _make_docx(tmp_path / "ua.docx", _p(), styles)
    units = resolve_effective(_first_run(view), view)["units"]
    assert units["size_pt"] == 11.0
    assert units["line_rule"] == "auto"
    assert (units["line_multiple"], units["line_pt"]) == (2.0, None)
    assert units["before_pt"] == 12.0
    assert (units["left_in"], units["first_line_in"]) == (1.0, 0.5)
    assert units["align"] == "left"            # "start" normalised


@pytest.mark.parametrize("rule", ["exact", "atLeast"])
def test_units_exact_line_spacing_in_points(tmp_path, rule):
    styles = _style("Normal", f'<w:pPr><w:spacing w:line="240" w:lineRule="{rule}"/></w:pPr>',
                    default=True)
    view = _make_docx(tmp_path / "ue.docx", _p(), styles)
    units = resolve_effective(_first_run(view), view)["units"]
    assert (units["line_rule"], units["line_pt"], units["line_multiple"]) == (rule, 12.0, None)


def test_units_hanging_indent_is_negative(tmp_path):
    styles = _style("Normal", "", default=True)
    view = _make_docx(tmp_path / "uh.docx", _p(ppr='<w:ind w:left="720" w:hanging="720"/>'), styles)
    units = resolve_effective(_first_run(view), view)["units"]
    assert (units["left_in"], units["first_line_in"]) == (0.5, -0.5)


def test_units_unreadable_number_is_none(tmp_path):
    styles = _style("Normal", '<w:pPr><w:spacing w:line="abc"/></w:pPr>'
                              '<w:rPr><w:sz w:val="big"/></w:rPr>', default=True)
    view = _make_docx(tmp_path / "ub.docx", _p(), styles)
    eff = resolve_effective(_first_run(view), view)
    assert eff["spacing"]["line"] == "abc"     # raw value kept as evidence
    assert eff["units"]["line_multiple"] is None
    assert eff["units"]["size_pt"] is None


# ---------------- on the K-0.3a fixture ----------------

def test_fixture_runs_resolve_to_apa_body_format(tmp_path):
    view = accepted_view(build_good_docx(tmp_path / "good.docx"))
    for p in view.paragraphs:
        for run in p.runs:
            eff = resolve_effective(run, view)
            assert (eff["font"], eff["size"]) == ("Calibri", "22"), p.text

    plain = [p for p in view.paragraphs if p.pPr is None and p.style_id is None and p.runs]
    assert len(plain) >= 5
    for p in plain:
        eff = resolve_effective(p.runs[0], view)
        assert eff["spacing"]["line"] == "480", p.text
        assert eff["spacing"]["lineRule"] == "auto"
        assert eff["indent"]["firstLine"] == "720"
        assert eff["align"] == "left"
        units = eff["units"]
        assert (units["size_pt"], units["line_multiple"]) == (11.0, 2.0)
        assert (units["first_line_in"], units["before_pt"], units["after_pt"]) == (0.5, 0.0, 0.0)
        assert units["align"] == "left"
