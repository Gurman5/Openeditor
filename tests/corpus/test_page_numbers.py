"""Tests for tools/corpus/checks/page_numbers.py (K-H2 step B: APA-05 page numbers).

The K-0.3a fixture has "different first page" on, a first-page header with a
w:fldSimple PAGE field and a default header with the fldChar/instrText form,
both right-aligned. Each test changes one thing.
Run:  python -m pytest tests/corpus/test_page_numbers.py -v
"""

import json

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from tools.corpus.checks.page_numbers import check_page_numbers, main
from tools.corpus.fixtures import _add_page_field_simple, build_good_docx


@pytest.fixture
def good(tmp_path):
    return build_good_docx(tmp_path / "good.docx")


def _edit(path, change):
    doc = Document(path)
    change(doc)
    doc.save(path)
    return path


def _sub(result):
    return result["subchecks"]["top_right_every_page"]


def _reasons(result):
    return [p["reason"] for p in _sub(result)["problems"]]


def _default_p(doc):
    return doc.sections[0].header.paragraphs[0]


def _first_p(doc):
    return doc.sections[0].first_page_header.paragraphs[0]


def _clear_runs(paragraph):
    for el in list(paragraph._p):
        if el.tag in (qn("w:r"), qn("w:fldSimple")):
            paragraph._p.remove(el)


def _tab_run():
    r = OxmlElement("w:r")
    r.append(OxmlElement("w:tab"))
    return r


def _instr_texts(doc):
    return list(_default_p(doc)._p.iter(qn("w:instrText")))


def test_fixture_passes(good):
    result = check_page_numbers(good)
    assert result["status"] == "pass", _sub(result)["problems"]
    assert result["matrix_id"] == "APA-05"
    forms = {s["type"]: s["page_field"] for s in _sub(result)["headers"]}
    assert forms == {"first": "fldSimple", "default": "fldChar"}


# ---------------- missing numbers ----------------

def test_first_page_header_without_number_fails(good):
    _edit(good, lambda d: _clear_runs(_first_p(d)))
    assert _reasons(check_page_numbers(good)) == ["first-page header has no page number"]


def test_title_page_with_no_first_page_header_fails(good):
    def change(doc):
        sect = doc.sections[0]._sectPr
        for ref in sect.findall(qn("w:headerReference")):
            if ref.get(qn("w:type")) == "first":
                sect.remove(ref)
    _edit(good, change)
    assert _reasons(check_page_numbers(good)) == ["no first-page header"]


def test_page_number_only_in_footer_fails(good):
    def change(doc):
        _clear_runs(_default_p(doc))
        _add_page_field_simple(doc.sections[0].footer.paragraphs[0])
    _edit(good, change)
    assert _reasons(check_page_numbers(good)) == [
        "default page number is in the footer, should be in the header"]


@pytest.mark.parametrize("instr, ok", [
    (" NUMPAGES ", False),
    (" PAGE \\* MERGEFORMAT ", True),
    (" page ", True),
])
def test_only_the_page_field_counts(good, instr, ok):
    def change(doc):
        _instr_texts(doc)[0].text = instr
    _edit(good, change)
    assert (check_page_numbers(good)["status"] == "pass") is ok


def test_instruction_split_over_runs_is_found(good):
    def change(doc):
        instr = _instr_texts(doc)[0]
        instr.text = " PA"
        rest = OxmlElement("w:r")
        rest_instr = OxmlElement("w:instrText")
        rest_instr.text = "GE "
        rest.append(rest_instr)
        instr.getparent().addnext(rest)
    _edit(good, change)
    assert check_page_numbers(good)["status"] == "pass"


# ---------------- position ----------------

def test_left_aligned_page_number_fails(good):
    _edit(good, lambda d: setattr(_default_p(d), "alignment", WD_ALIGN_PARAGRAPH.LEFT))
    result = check_page_numbers(good)
    assert _reasons(result) == ["default page number is not at the right"]


@pytest.mark.parametrize("tabs, ok", [(2, True), (1, False)])
def test_right_tab_stop_from_header_style(good, tabs, ok):
    # Word's Header style: center stop at 4680, right stop at 9360.
    def change(doc):
        p = _default_p(doc)
        p.alignment = None                      # no right alignment any more
        for _ in range(tabs):
            p._p.insert(1, _tab_run())          # index 0 is w:pPr
    _edit(good, change)
    result = check_page_numbers(good)
    assert (result["status"] == "pass") is ok, _sub(result)
    if ok:
        slot = next(s for s in _sub(result)["headers"] if s["type"] == "default")
        assert slot["position"] == "right tab stop at 9360 twips"


# ---------------- even/odd + sections ----------------

def test_even_and_odd_without_even_header_fails(good):
    _edit(good, lambda d: setattr(d.settings, "odd_and_even_pages_header_footer", True))
    result = check_page_numbers(good)
    assert _sub(result)["even_and_odd_headers"] is True
    assert _reasons(result) == ["no even-page header"]


def test_even_and_odd_with_even_header_passes(good):
    def change(doc):
        doc.settings.odd_and_even_pages_header_footer = True
        p = doc.sections[0].even_page_header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _add_page_field_simple(p)
    _edit(good, change)
    assert check_page_numbers(good)["status"] == "pass"


def test_second_section_inherits_headers(good):
    _edit(good, lambda d: d.add_section())
    result = check_page_numbers(good)
    assert result["status"] == "pass", _sub(result)["problems"]
    assert {s["section"] for s in _sub(result)["headers"]} == {0, 1}


def test_second_section_with_own_empty_header_fails(good):
    def change(doc):
        doc.add_section().header.is_linked_to_previous = False
    _edit(good, change)
    problems = _sub(check_page_numbers(good))["problems"]
    assert [(p["section"], p["reason"]) for p in problems] == [
        (1, "default header has no page number")]


def test_cli_prints_json(good, capsys):
    assert main(["--docx", str(good)]) == 0
    assert json.loads(capsys.readouterr().out)["status"] == "pass"
