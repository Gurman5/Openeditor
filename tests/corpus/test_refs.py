"""Tests for tools/corpus/checks/refs.py (K-H5, APA-12 references + APA-18 DOIs).

The K-0.3a fixture has a bold, centered "References" label with a direct
pageBreakBefore, then three entries (Adams, Chen, Davis) with hanging indents
and https://doi.org/ DOIs. Each test changes one thing.
Run:  python -m pytest tests/corpus/test_refs.py -v
"""

import copy
import json

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.text.paragraph import Paragraph

from tools.corpus.checks.refs import check_references, main
from tools.corpus.fixtures import build_good_docx


@pytest.fixture
def good(tmp_path):
    return build_good_docx(tmp_path / "good.docx")


def _edit(path, change):
    """Open with python-docx, call change(doc, label, entries), save in place."""
    doc = Document(path)
    i = next(i for i, p in enumerate(doc.paragraphs) if p.text == "References")
    change(doc, doc.paragraphs[i], doc.paragraphs[i + 1:i + 4])
    doc.save(path)
    return path


def _sub(result, row, name):
    return result[row]["subchecks"][name]


def _status(result, row, name=None):
    return result[row]["status"] if name is None else _sub(result, row, name)["status"]


def _no_break(label):
    label.paragraph_format.page_break_before = None


def _previous(doc, label):
    return doc.paragraphs[[p._p for p in doc.paragraphs].index(label._p) - 1]


def _sect_pr(continuous=False):
    sect = OxmlElement("w:sectPr")
    if continuous:
        kind = OxmlElement("w:type")
        kind.set(qn("w:val"), "continuous")
        sect.append(kind)
    return sect


def test_fixture_passes_both_rows(good):
    result = check_references(good)
    assert _status(result, "APA-12") == "pass", result["APA-12"]
    assert _status(result, "APA-18") == "pass", result["APA-18"]
    assert _sub(result, "APA-12", "new_page")["how"] == "pageBreakBefore on the label"
    assert _sub(result, "APA-12", "entries_present")["entries"] == 3
    assert _sub(result, "APA-18", "https_doi_org")["dois_found"] == 3
    assert _sub(result, "APA-12", "close_calls")["items"] == []


# ---------------- new page: all three ways ----------------

def test_no_page_break_fails(good):
    _edit(good, lambda d, label, entries: _no_break(label))
    result = check_references(good)
    assert _status(result, "APA-12", "new_page") == "fail"
    assert _status(result, "APA-12") == "fail"


def test_page_break_before_through_style_passes(good):
    def change(doc, label, entries):
        style = doc.styles.add_style("NewPageLabel", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Normal"]
        style.paragraph_format.page_break_before = True
        label.style = style
        _no_break(label)
    _edit(good, change)
    sub = _sub(check_references(good), "APA-12", "new_page")
    assert (sub["status"], sub["how"]) == ("pass", "pageBreakBefore on the label")


def test_page_break_at_end_of_previous_paragraph_passes(good):
    def change(doc, label, entries):
        _no_break(label)
        _previous(doc, label).add_run().add_break(WD_BREAK.PAGE)
    _edit(good, change)
    sub = _sub(check_references(good), "APA-12", "new_page")
    assert sub["status"] == "pass"
    assert sub["how"].startswith("page break after the text of paragraph")


def test_page_break_in_empty_paragraph_passes(good):
    def change(doc, label, entries):
        _no_break(label)
        empty = OxmlElement("w:p")
        label._p.addprevious(empty)
        Paragraph(empty, label._parent).add_run().add_break(WD_BREAK.PAGE)
    _edit(good, change)
    sub = _sub(check_references(good), "APA-12", "new_page")
    assert sub["status"] == "pass"
    assert sub["how"].startswith("page break in empty paragraph")


def test_page_break_at_start_of_label_passes(good):
    def change(doc, label, entries):
        _no_break(label)
        run = OxmlElement("w:r")
        br = OxmlElement("w:br")
        br.set(qn("w:type"), "page")
        run.append(br)
        label._p.insert(1, run)              # index 0 is w:pPr
    _edit(good, change)
    sub = _sub(check_references(good), "APA-12", "new_page")
    assert (sub["status"], sub["how"]) == ("pass", "page break at the start of the label paragraph")


@pytest.mark.parametrize("continuous, ok", [(False, True), (True, False)])
def test_section_break_counts_unless_continuous(good, continuous, ok):
    def change(doc, label, entries):
        _no_break(label)
        _previous(doc, label)._p.get_or_add_pPr().append(_sect_pr(continuous))
    _edit(good, change)
    assert (_status(check_references(good), "APA-12", "new_page") == "pass") is ok


# ---------------- label + entry formatting ----------------

def test_label_not_bold_fails(good):
    _edit(good, lambda d, label, entries: setattr(label.runs[0], "bold", False))
    assert _status(check_references(good), "APA-12", "label_bold") == "fail"


def test_label_not_centered_fails(good):
    _edit(good, lambda d, label, entries: setattr(label, "alignment", WD_ALIGN_PARAGRAPH.LEFT))
    assert _status(check_references(good), "APA-12", "label_centered") == "fail"


def test_entry_without_hanging_indent_fails(good):
    def change(doc, label, entries):
        pf = entries[1].paragraph_format
        pf.left_indent, pf.first_line_indent = Inches(0), Inches(0)
    _edit(good, change)
    problems = _sub(check_references(good), "APA-12", "hanging_indent")["problems"]
    assert [p["entry"][:5] for p in problems] == ["Chen,"]


def test_single_spaced_entry_fails(good):
    _edit(good, lambda d, label, entries: setattr(entries[0].paragraph_format, "line_spacing", 1.0))
    assert _status(check_references(good), "APA-12", "double_spacing") == "fail"


# ---------------- alphabetical order ----------------

def test_out_of_order_entries_fail(good):
    _edit(good, lambda d, label, entries: entries[2]._p.addnext(entries[0]._p))  # Adams to end
    sub = _sub(check_references(good), "APA-12", "alphabetical")
    assert sub["status"] == "fail"
    assert sub["problems"][0]["keys"] == ["davis", "adams"]


def test_tie_is_info_only(good):
    _edit(good, lambda d, label, entries: entries[0]._p.addnext(copy.deepcopy(entries[0]._p)))
    result = check_references(good)
    items = _sub(result, "APA-12", "close_calls")["items"]
    assert items[0]["reason"].startswith("same first word")
    assert _status(result, "APA-12", "alphabetical") == "pass"


def test_article_start_is_info_only(good):
    def change(doc, label, entries):
        new = copy.deepcopy(entries[2]._p)
        entries[2]._p.addnext(new)
        para = Paragraph(new, entries[2]._parent)
        for run in para.runs[1:]:
            run.text = ""
        para.runs[0].text = "The Example Group. (2022). Annual report. https://doi.org/10.0000/x.1"
    _edit(good, change)
    result = check_references(good)
    reasons = [i["reason"] for i in _sub(result, "APA-12", "close_calls")["items"]]
    assert any(r.startswith('starts with "The"') for r in reasons)
    assert _status(result, "APA-12") == "pass"


# ---------------- APA-18 DOIs ----------------

GOOD_DOI = "https://doi.org/10.0000/jep.2019.003"


def _set_adams_doi(entries, form):
    entries[0].runs[2].text = f"(3), 45–60. {form}"


@pytest.mark.parametrize("form, reason", [
    ("doi:10.0000/jep.2019.003", "doi: prefix"),
    ("DOI: 10.0000/jep.2019.003", "doi: prefix"),
    ("http://dx.doi.org/10.0000/jep.2019.003", "dx.doi.org"),
    ("https://dx.doi.org/10.0000/jep.2019.003", "dx.doi.org"),
    ("http://doi.org/10.0000/jep.2019.003", "http:// instead of https://"),
    ("10.0000/jep.2019.003", "bare DOI without https://doi.org/"),
])
def test_old_doi_forms_fail(good, form, reason):
    _edit(good, lambda d, label, entries: _set_adams_doi(entries, form))
    sub = _sub(check_references(good), "APA-18", "https_doi_org")
    assert sub["status"] == "fail"
    assert [p["reason"] for p in sub["problems"]] == [reason]


def test_entry_without_doi_is_fine(good):
    _edit(good, lambda d, label, entries: _set_adams_doi(entries, ""))
    result = check_references(good)
    assert _status(result, "APA-18") == "pass"
    assert _sub(result, "APA-18", "https_doi_org")["dois_found"] == 2


def test_no_dois_at_all_is_na(good):
    def change(doc, label, entries):
        for p in entries:
            for run in p.runs:
                run.text = run.text.split("https://")[0]
    _edit(good, change)
    row = check_references(good)["APA-18"]
    assert (row["status"], row["reason"]) == ("na", "no DOIs in the reference list")


# ---------------- list boundaries ----------------

def test_list_ends_at_next_label(good):
    def change(doc, label, entries):
        appendix = OxmlElement("w:p")
        entries[2]._p.addnext(appendix)
        para = Paragraph(appendix, entries[2]._parent)
        para.add_run("Appendix").bold = True
        text = OxmlElement("w:p")
        appendix.addnext(text)
        Paragraph(text, entries[2]._parent).add_run("Appendix text, not an entry.")
    _edit(good, change)
    result = check_references(good)
    assert _sub(result, "APA-12", "entries_present")["entries"] == 3
    assert _status(result, "APA-12") == "pass"


def test_no_references_label_is_na(good):
    _edit(good, lambda d, label, entries: label._p.getparent().remove(label._p))
    result = check_references(good)
    assert (_status(result, "APA-12"), _status(result, "APA-18")) == ("na", "na")


def test_label_with_no_entries_fails(good):
    def change(doc, label, entries):
        for p in entries:
            p._p.getparent().remove(p._p)
    _edit(good, change)
    result = check_references(good)
    assert _status(result, "APA-12", "entries_present") == "fail"
    assert _status(result, "APA-18") == "na"


def test_cli_prints_json(good, capsys):
    assert main(["--docx", str(good)]) == 0
    assert json.loads(capsys.readouterr().out)["APA-12"]["status"] == "pass"
