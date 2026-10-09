from pathlib import Path
import zipfile

import pytest
from docx import Document
from lxml import etree


CORPUS_ROOT = Path(__file__).parent

MAX_FILE_SIZE_MB = 20
MAX_FILE_SIZE_BYTES = MAX_FILE_SIZE_MB * 1024 * 1024
MAX_WORD_COUNT = 15_000

FORBIDDEN_EXACT = {
    "word/vbaProject.bin",
}

FORBIDDEN_PREFIXES = (
    "word/embeddings/",
)


def _find_docx_files():
    """Return every DOCX file currently stored in the Track K corpus."""
    return sorted(CORPUS_ROOT.rglob("*.docx"))


DOCX_FILES = _find_docx_files()

SEED_DIR = CORPUS_ROOT / "seed"
SEED_FILES = sorted(SEED_DIR.glob("*.docx"))


def test_seed_corpus_has_exactly_four_docx_files():
    assert len(SEED_FILES) == 4, (
        f"Expected exactly 4 seed DOCX files in {SEED_DIR}, "
        f"but found {len(SEED_FILES)}"
    )


def _doc_id(path: Path) -> str:
    """Readable pytest ID for each corpus document."""
    return str(path.relative_to(CORPUS_ROOT))

def _doc_id(path: Path) -> str:
    """Readable pytest ID for each corpus document."""
    return str(path.relative_to(CORPUS_ROOT))


def _count_words(docx_path: Path) -> int:
    """Count words in normal document paragraphs using python-docx."""
    document = Document(docx_path)

    words = 0

    for paragraph in document.paragraphs:
        words += len(paragraph.text.split())

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    words += len(paragraph.text.split())

    return words


@pytest.mark.parametrize(
    "docx_path",
    DOCX_FILES,
    ids=_doc_id,
)
def test_corpus_docx_passes_upload_gate(docx_path: Path):
    """
    Every corpus DOCX must pass the basic OpenEditor upload requirements.

    Checks:
    - valid ZIP/DOCX container
    - XML files are well formed
    - no VBA macro project
    - no embedded objects
    - file size <= 20 MB
    - word count <= 15,000
    """

    # 1. File size
    file_size = docx_path.stat().st_size

    assert file_size <= MAX_FILE_SIZE_BYTES, (
        f"{docx_path.name} is too large: "
        f"{file_size / (1024 * 1024):.2f} MB "
        f"(maximum {MAX_FILE_SIZE_MB} MB)"
    )

    # 2. Valid ZIP container
    assert zipfile.is_zipfile(docx_path), (
        f"{docx_path.name} is not a valid DOCX/ZIP file"
    )

    with zipfile.ZipFile(docx_path, "r") as archive:
        members = archive.namelist()

        # 3. Check ZIP integrity
        bad_member = archive.testzip()

        assert bad_member is None, (
            f"{docx_path.name} contains a corrupt ZIP member: {bad_member}"
        )

        # 4. No VBA macro project
        for forbidden in FORBIDDEN_EXACT:
            assert forbidden not in members, (
                f"{docx_path.name} contains forbidden file: {forbidden}"
            )

        # 5. No embedded files
        embedded_files = [
            member
            for member in members
            if member.startswith(FORBIDDEN_PREFIXES)
            and not member.endswith("/")
        ]

        assert not embedded_files, (
            f"{docx_path.name} contains embedded files: "
            f"{embedded_files}"
        )

        # 6. XML must be well formed
        xml_members = [
            member
            for member in members
            if member.endswith(".xml")
        ]

        for member in xml_members:
            xml_data = archive.read(member)

            try:
                etree.fromstring(xml_data)
            except etree.XMLSyntaxError as exc:
                pytest.fail(
                    f"{docx_path.name}: invalid XML in "
                    f"{member}: {exc}"
                )

    # 7. Word count
    try:
        word_count = _count_words(docx_path)
    except Exception as exc:
        pytest.fail(
            f"{docx_path.name} could not be opened "
            f"with python-docx: {exc}"
        )

    assert word_count <= MAX_WORD_COUNT, (
        f"{docx_path.name} contains {word_count:,} words "
        f"(maximum {MAX_WORD_COUNT:,})"
    )