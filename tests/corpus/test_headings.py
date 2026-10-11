"""Tests for tools/corpus/checks/headings.py (K-H4, APA-10 introduction + APA-11 headings).

Each test starts from the K-0.3a synthetic fixture and changes one thing.
Run:  python -m pytest tests/corpus/test_headings.py -v
"""

import json

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.text.paragraph import Paragraph

from tools.corpus.checks.headings import check_headings, main
from tools.corpus.fixtures import TITLE, build_good_docx


@pytest.fixture
def good(tmp_path):
    return build_good_docx(tmp_path / "good.docx")


def _edit(path, change):
    doc = Document(path)
    change(doc)
    doc.save(path)
    return path


def _find(doc, text):
    return next(p for p in doc.paragraphs if p.text == text)


def _insert_after(paragraph, runs, style=None):
    """New paragraph after `paragraph`. runs = [(text, bold, italic), ...]."""
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    para = Paragraph(new_p, paragraph._parent)
    if style:
        para.style = style
    for text, bold, italic in runs:
        run = para.add_run(text)
        run.bold, run.italic = bold or None, italic or None
    return para


def _headings(result):
    return [(h["level"], h["text"], h["status"]) for h in result["APA-11"]["headings"]]


def _problems(result, text):
    return next(h["problems"] for h in result["APA-11"]["headings"] if h["text"] == text)


BODY = (" Students were invited by email.", False, False)


def test_fixture_passes_both_rows(good):
    result = check_headings(good)
    assert result["APA-10"]["status"] == "pass", result["APA-10"]
    assert result["APA-11"]["status"] == "pass", result["APA-11"]
    assert _headings(result) == [(1, "Method", "pass"), (2, "Participants", "pass"),
                                 (2, "Procedure", "pass"), (1, "Results", "pass"),
                                 (1, "Discussion", "pass")]
    assert result["APA-11"]["title_case"]["headings"] == []


# ---------------- APA-10 ----------------

def test_introduction_heading_fails(good):
    _edit(good, lambda d: _insert_after(_find(d, "Method"), [("Introduction", False, False)],
                                        style="Heading 1"))
    result = check_headings(good)
    assert result["APA-10"]["status"] == "fail"
    assert result["APA-10"]["subchecks"]["no_introduction_heading"]["status"] == "fail"


def test_introduction_run_in_heading_fails(good):
    _edit(good, lambda d: _insert_after(_find(d, "Participants"),
                                        [("Introduction.", True, False), BODY]))
    assert check_headings(good)["APA-10"]["subchecks"]["no_introduction_heading"]["status"] == "fail"


def test_missing_repeated_title_fails(good):
    def change(doc):
        copy = [p for p in doc.paragraphs if p.text == TITLE][1]
        copy._p.getparent().remove(copy._p)
    _edit(good, change)
    sub = check_headings(good)["APA-10"]["subchecks"]["title_repeated"]
    assert (sub["status"], sub["copies"]) == ("fail", [])


def test_repeated_title_not_bold_fails(good):
    def change(doc):
        [p for p in doc.paragraphs if p.text == TITLE][1].runs[0].bold = False
    _edit(good, change)
    assert check_headings(good)["APA-10"]["subchecks"]["title_repeated"]["status"] == "fail"


# ---------------- APA-11: levels 1-3 ----------------

def test_level1_not_centered_fails(good):
    _edit(good, lambda d: setattr(_find(d, "Method"), "alignment", WD_ALIGN_PARAGRAPH.LEFT))
    result = check_headings(good)
    assert result["APA-11"]["status"] == "fail"
    assert _problems(result, "Method") == ["alignment should be center, got left"]


def test_level2_centered_fails(good):
    _edit(good, lambda d: setattr(_find(d, "Participants"), "alignment", WD_ALIGN_PARAGRAPH.CENTER))
    assert _problems(check_headings(good), "Participants") == ["alignment should be left, got center"]


def _add_level3(doc, italic=True):
    para = _insert_after(_find(doc, "Participants"), [("Sampling Approach", True, italic)],
                         style="Heading 3")
    para.paragraph_format.first_line_indent = Inches(0)


def test_level3_bold_italic_passes(good):
    _edit(good, _add_level3)
    result = check_headings(good)
    assert (3, "Sampling Approach", "pass") in _headings(result)


def test_level3_not_italic_fails(good):
    _edit(good, lambda d: _add_level3(d, italic=False))
    assert _problems(check_headings(good), "Sampling Approach") == ["should be bold italic"]


# ---------------- APA-11: levels 4-5 (run-in) ----------------

def test_level4_run_in_passes(good):
    _edit(good, lambda d: _insert_after(_find(d, "Participants"),
                                        [("Recruitment Method.", True, False), BODY]))
    heading = next(h for h in check_headings(good)["APA-11"]["headings"]
                   if h["text"] == "Recruitment Method.")
    assert (heading["level"], heading["found_by"], heading["status"]) == (4, "run-in", "pass")


def test_level5_run_in_passes(good):
    _edit(good, lambda d: _insert_after(_find(d, "Participants"),
                                        [("Email Invitations.", True, True), BODY]))
    assert (5, "Email Invitations.", "pass") in _headings(check_headings(good))


def test_run_in_heading_not_indented_fails(good):
    def change(doc):
        para = _insert_after(_find(doc, "Participants"), [("Recruitment Method.", True, False), BODY])
        para.paragraph_format.first_line_indent = Inches(0)
    _edit(good, change)
    assert _problems(check_headings(good), "Recruitment Method.") == [
        "first-line indent should be 0.5in, got 0.0"]


def test_level4_on_its_own_line_fails(good):
    _edit(good, lambda d: _insert_after(_find(d, "Participants"),
                                        [("Recruitment Method", True, False)], style="Heading 4"))
    problems = _problems(check_headings(good), "Recruitment Method")
    assert "Level 4 headings should run in with the paragraph text" in problems


# ---------------- numbered headings ----------------

@pytest.mark.parametrize("text, found", [
    ("1. Method", '"1."'),
    ("2.1 Method", '"2.1"'),
    ("IV. Method", '"IV."'),
    ("A. Method", '"A."'),
])
def test_typed_heading_number_fails(good, text, found):
    _edit(good, lambda d: setattr(_find(d, "Method").runs[0], "text", text))
    problems = _problems(check_headings(good), text)
    assert problems == [f"heading should not be numbered (typed {found})"]


def test_year_at_start_of_heading_is_not_numbering(good):
    _edit(good, lambda d: setattr(_find(d, "Results").runs[0], "text", "2020 Survey Results"))
    assert check_headings(good)["APA-11"]["status"] == "pass"


def _num_pr(num_id="1"):
    num = OxmlElement("w:numPr")
    ilvl, nid = OxmlElement("w:ilvl"), OxmlElement("w:numId")
    ilvl.set(qn("w:val"), "0")
    nid.set(qn("w:val"), num_id)
    num.extend([ilvl, nid])
    return num


def test_automatic_numbering_on_paragraph_fails(good):
    _edit(good, lambda d: _find(d, "Method")._p.get_or_add_pPr().insert(1, _num_pr()))
    assert _problems(check_headings(good), "Method") == [
        "heading should not be numbered (automatic list numbering from direct)"]


def test_automatic_numbering_from_style_fails(good):
    def change(doc):
        style_ppr = doc.styles["Heading 1"].element.get_or_add_pPr()
        style_ppr.insert(0, _num_pr())
    _edit(good, change)
    problems = _problems(check_headings(good), "Method")
    assert problems == ["heading should not be numbered (automatic list numbering from "
                        "style:Heading1)"]


def test_numbering_switched_off_with_num_id_zero_passes(good):
    _edit(good, lambda d: _find(d, "Method")._p.get_or_add_pPr().insert(1, _num_pr("0")))
    assert check_headings(good)["APA-11"]["status"] == "pass"


# ---------------- not headings ----------------

def test_bold_word_without_period_is_not_a_heading(good):
    _edit(good, lambda d: _insert_after(_find(d, "Participants"),
                                        [("Important", True, False), BODY]))
    assert len(check_headings(good)["APA-11"]["headings"]) == 5


@pytest.mark.parametrize("label", ["Abstract", "Appendix A", "Author Note"])
def test_section_labels_with_heading_style_are_not_headings(good, label):
    def change(doc):
        if label == "Abstract":
            _find(doc, "Abstract").style = "Heading 1"
        else:
            _insert_after(_find(doc, "Discussion"), [(label, True, False)], style="Heading 1")
    _edit(good, change)
    result = check_headings(good)
    assert len(result["APA-11"]["headings"]) == 5
    assert result["APA-11"]["status"] == "pass"


# ---------------- title case + edge cases ----------------

def test_title_case_is_info_only(good):
    _edit(good, lambda d: setattr(_find(d, "Results").runs[0], "text", "Results of the study"))
    result = check_headings(good)
    info = result["APA-11"]["title_case"]
    assert info["status"] == "info"
    assert info["headings"][0]["words"] == ["study"]
    assert result["APA-11"]["status"] == "pass"


def test_no_headings_is_na(tmp_path):
    path = tmp_path / "plain.docx"
    doc = Document()
    doc.add_paragraph("A Title")
    doc.add_paragraph("Some body text.")
    doc.save(path)
    assert check_headings(path)["APA-11"]["status"] == "na"


def test_cli_prints_json(good, capsys):
    assert main(["--docx", str(good)]) == 0
    result = json.loads(capsys.readouterr().out)
    assert result["APA-11"]["status"] == "pass"
