"""Tests for tools/corpus/checks/abstract.py (K-H6, APA-08 abstract + APA-09 keywords).

Each test starts from the K-0.3a synthetic fixture and changes one thing.
Run:  python -m pytest tests/corpus/test_abstract.py -v
"""

import copy
import json

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.shared import Inches

from tools.corpus.checks.abstract import ENV_VAR, check_abstract, main
from tools.corpus.fixtures import TITLE, build_good_docx


@pytest.fixture
def good(tmp_path):
    return build_good_docx(tmp_path / "good.docx")


def _edit(path, change):
    """Open with python-docx, call change(doc, label, body, keywords), save in place."""
    doc = Document(path)
    i = next(i for i, p in enumerate(doc.paragraphs) if p.text == "Abstract")
    change(doc, *doc.paragraphs[i:i + 3])
    doc.save(path)
    return path


def _remove(paragraph):
    paragraph._p.getparent().remove(paragraph._p)


def _status(result, row, sub=None):
    return result[row]["status"] if sub is None else result[row]["subchecks"][sub]["status"]


def test_fixture_passes_both_rows(good):
    result = check_abstract(good)
    assert _status(result, "APA-08") == "pass", result["APA-08"]
    assert _status(result, "APA-09") == "pass", result["APA-09"]
    assert result["APA-08"]["subchecks"]["word_count"]["words"] < 250
    assert result["APA-09"]["subchecks"]["capitalised_words"]["words"] == []


# ---------------- APA-08 ----------------

def test_abstract_over_250_words_fails(good):
    _edit(good, lambda d, label, body, kw: body.add_run(" extra" * 251))
    result = check_abstract(good)
    assert _status(result, "APA-08", "word_count") == "fail"
    assert _status(result, "APA-08") == "fail"


def test_abstract_two_paragraphs_fails(good):
    _edit(good, lambda d, label, body, kw: body._p.addnext(copy.deepcopy(body._p)))
    result = check_abstract(good)
    assert _status(result, "APA-08", "one_paragraph") == "fail"
    assert len(result["APA-08"]["subchecks"]["one_paragraph"]["paragraphs"]) == 2


def test_abstract_indented_fails(good):
    _edit(good, lambda d, label, body, kw: setattr(body.paragraph_format,
                                                   "first_line_indent", Inches(0.5)))
    result = check_abstract(good)
    assert _status(result, "APA-08", "unindented") == "fail"
    assert result["APA-08"]["subchecks"]["unindented"]["indented"][0]["first_line_in"] == 0.5


def test_label_not_bold_fails(good):
    _edit(good, lambda d, label, body, kw: setattr(label.runs[0], "bold", False))
    assert _status(check_abstract(good), "APA-08", "label_bold") == "fail"


def test_label_not_centered_fails(good):
    _edit(good, lambda d, label, body, kw: setattr(label, "alignment", WD_ALIGN_PARAGRAPH.LEFT))
    result = check_abstract(good)
    assert _status(result, "APA-08", "label_centered") == "fail"
    assert result["APA-08"]["subchecks"]["label_centered"]["align"] == "left"


def test_label_bold_centered_through_a_style_passes(good):
    def change(doc, label, body, kw):
        style = doc.styles.add_style("SectionLabel", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Normal"]
        style.font.bold = True
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style.paragraph_format.first_line_indent = Inches(0)
        label.style = style
        label.alignment = None          # remove direct centering
        label.runs[0].bold = None       # remove direct bold
    _edit(good, change)
    assert _status(check_abstract(good), "APA-08") == "pass"


# ---------------- APA-09 ----------------

def test_keywords_label_not_italic_fails(good):
    _edit(good, lambda d, label, body, kw: setattr(kw.runs[0], "italic", False))
    result = check_abstract(good)
    assert _status(result, "APA-09", "label_italic") == "fail"
    assert result["APA-09"]["subchecks"]["label_italic"]["not_italic_runs"] == ["Keywords:"]


def test_keywords_not_indented_fails(good):
    _edit(good, lambda d, label, body, kw: setattr(kw.paragraph_format,
                                                   "first_line_indent", Inches(0)))
    assert _status(check_abstract(good), "APA-09", "indented") == "fail"


def test_keywords_too_long_fails(good):
    _edit(good, lambda d, label, body, kw: kw.add_run(", another keyword" * 6))
    result = check_abstract(good)
    assert _status(result, "APA-09", "one_line") == "fail"
    assert result["APA-09"]["subchecks"]["one_line"]["max"] == 100


def test_keyword_limit_can_be_changed(good, monkeypatch):
    _edit(good, lambda d, label, body, kw: kw.add_run(", another keyword" * 6))
    assert _status(check_abstract(good, max_keyword_chars=200), "APA-09") == "pass"
    monkeypatch.setenv(ENV_VAR, "200")
    assert _status(check_abstract(good), "APA-09") == "pass"


def test_capitalised_keyword_is_info_only(good):
    def change(doc, label, body, kw):
        kw.runs[1].text = " spaced practice, Australia, Retention"
    _edit(good, change)
    result = check_abstract(good)
    sub = result["APA-09"]["subchecks"]["capitalised_words"]
    assert (sub["status"], sub["words"]) == ("info", ["Australia", "Retention"])
    assert _status(result, "APA-09") == "pass"


# ---------------- missing parts + page breaks ----------------

def test_no_abstract_is_na(good):
    def change(doc, label, body, kw):
        for p in (label, body, kw):
            _remove(p)
    _edit(good, change)
    result = check_abstract(good)
    assert (_status(result, "APA-08"), _status(result, "APA-09")) == ("na", "na")


def test_no_keywords_is_na_for_apa09_only(good):
    _edit(good, lambda d, label, body, kw: _remove(kw))
    result = check_abstract(good)
    assert (_status(result, "APA-08"), _status(result, "APA-09")) == ("pass", "na")


# Boundary tests: the keywords line is removed so the abstract area has to end
# some other way. In the fixture the body title right after the abstract has a
# direct pageBreakBefore; _no_boundary() removes it too.

def _body_title(doc):
    return next(p for p in doc.paragraphs
                if p.text == TITLE and p.paragraph_format.page_break_before)


def _no_boundary(doc, kw):
    _remove(kw)
    title = _body_title(doc)
    title.paragraph_format.page_break_before = None
    return title


def _abstract_paragraphs(result):
    return result["APA-08"]["subchecks"]["one_paragraph"]["paragraphs"]


def test_without_a_boundary_body_text_is_swept_in(good):
    # Control: proves the boundary tests below really depend on their boundary.
    _edit(good, lambda d, label, body, kw: _no_boundary(d, kw))
    result = check_abstract(good)
    assert _status(result, "APA-08", "one_paragraph") == "fail"
    assert len(_abstract_paragraphs(result)) > 1


def test_direct_page_break_before_ends_the_abstract(good):
    _edit(good, lambda d, label, body, kw: _remove(kw))
    result = check_abstract(good)
    assert _status(result, "APA-08", "one_paragraph") == "pass", result["APA-08"]


def test_style_page_break_before_ends_the_abstract(good):
    def change(doc, label, body, kw):
        title = _no_boundary(doc, kw)
        style = doc.styles.add_style("NewPageTitle", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Normal"]
        style.paragraph_format.page_break_before = True
        title.style = style
    _edit(good, change)
    result = check_abstract(good)
    assert _status(result, "APA-08", "one_paragraph") == "pass", result["APA-08"]


def test_page_break_inside_a_run_ends_the_abstract(good):
    def change(doc, label, body, kw):
        _no_boundary(doc, kw)
        body.add_run().add_break(WD_BREAK.PAGE)
    _edit(good, change)
    result = check_abstract(good)
    assert _status(result, "APA-08", "one_paragraph") == "pass", result["APA-08"]


def test_section_break_ends_the_abstract(good):
    def change(doc, label, body, kw):
        _no_boundary(doc, kw)
        body._p.get_or_add_pPr().append(OxmlElement("w:sectPr"))
    _edit(good, change)
    result = check_abstract(good)
    assert _status(result, "APA-08", "one_paragraph") == "pass", result["APA-08"]
    assert len(_abstract_paragraphs(result)) == 1   # the sectPr paragraph itself counts


def test_cli_prints_json(good, capsys):
    assert main(["--docx", str(good)]) == 0
    result = json.loads(capsys.readouterr().out)
    assert result["APA-08"]["status"] == "pass"
