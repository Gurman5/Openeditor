"""Tests for tools/corpus/checks/layout.py (K-H2 step A: APA-00/01/03/04).

Each test starts from the K-0.3a synthetic fixture and changes one thing.
Run:  python -m pytest tests/corpus/test_layout.py -v
"""

import json

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, Twips
from docx.text.paragraph import Paragraph

from tools.corpus.checks.layout import (
    CHECKS, MARGIN_ENV_VAR, check_alignment, check_indent, check_layout, check_margins,
    check_spacing, main,
)
from tools.corpus.fixtures import build_good_docx


@pytest.fixture
def good(tmp_path):
    return build_good_docx(tmp_path / "good.docx")


def _edit(path, change):
    doc = Document(path)
    change(doc)
    doc.save(path)
    return path


def _find(doc, start):
    return next(p for p in doc.paragraphs if p.text.startswith(start))


def _new_paragraph(anchor, text="", before=False):
    new_p = OxmlElement("w:p")
    (anchor._p.addprevious if before else anchor._p.addnext)(new_p)
    para = Paragraph(new_p, anchor._parent)
    if text:
        para.add_run(text)
    return para


def _sub(result, name):
    return result["subchecks"][name]


BODY = "Forty students"          # a plain body paragraph in the fixture


def test_fixture_passes_all_layout_checks(good):
    result = check_layout(good)
    assert set(result) == {"APA-00", "APA-01", "APA-03", "APA-04"}
    for row in result.values():
        assert row["status"] == "pass", row


# ---------------- APA-00 margins ----------------

def test_wrong_margin_fails(good):
    _edit(good, lambda d: setattr(d.sections[0], "left_margin", Inches(0.75)))
    sub = _sub(check_margins(good), "one_inch")
    assert sub["status"] == "fail"
    assert sub["problems"] == [{"section": 0, "side": "left", "twips": "1080", "inches": 0.75}]


def test_wrong_margin_in_second_section_fails(good):
    def change(doc):
        doc.add_section().top_margin = Inches(1.25)
    _edit(good, change)
    sub = _sub(check_margins(good), "one_inch")
    assert [(p["section"], p["side"]) for p in sub["problems"]] == [(1, "top")]


def test_margin_tolerance_setting(good, monkeypatch):
    _edit(good, lambda d: setattr(d.sections[0], "left_margin", Twips(1417)))   # 2.5 cm
    assert check_margins(good)["status"] == "fail"                              # default 0
    assert check_margins(good, tolerance_twips=30)["status"] == "pass"
    monkeypatch.setenv(MARGIN_ENV_VAR, "30")
    assert check_margins(good)["status"] == "pass"


# ---------------- APA-01 spacing ----------------

def test_one_point_one_five_spacing_fails(good):
    _edit(good, lambda d: setattr(d.styles["Normal"].paragraph_format, "line_spacing", 1.15))
    sub = _sub(check_spacing(good), "double_spacing")
    assert sub["status"] == "fail"
    assert sub["problems"][0]["line_multiple"] == 1.15


def test_exact_24pt_spacing_fails(good):
    _edit(good, lambda d: setattr(d.styles["Normal"].paragraph_format, "line_spacing", Pt(24)))
    problem = _sub(check_spacing(good), "double_spacing")["problems"][0]
    assert (problem["line_rule"], problem["line_pt"]) == ("exact", 24.0)


def test_space_after_is_its_own_subcheck(good):
    _edit(good, lambda d: setattr(d.styles["Normal"].paragraph_format, "space_after", Pt(8)))
    result = check_spacing(good)
    assert _sub(result, "no_space_before_after")["status"] == "fail"
    assert _sub(result, "no_space_before_after")["problems"][0]["after_pt"] == 8.0
    assert _sub(result, "double_spacing")["status"] == "pass"


@pytest.mark.parametrize("before", [False, True])
def test_blank_line_next_to_heading_fails(good, before):
    _edit(good, lambda d: _new_paragraph(_find(d, "Method"), before=before))
    sub = _sub(check_spacing(good), "no_blank_lines_around_headings")
    assert sub["status"] == "fail"
    assert sub["problems"][0]["heading"] == "Method"
    assert sub["problems"][0]["where"] == ("before" if before else "after")


def test_blank_line_on_title_page_is_allowed(good):
    _edit(good, lambda d: _new_paragraph(d.paragraphs[0], before=True))
    assert check_spacing(good)["status"] == "pass"


def test_single_spaced_references_fail(good):
    def change(doc):
        start = next(i for i, p in enumerate(doc.paragraphs) if p.text == "References")
        for p in doc.paragraphs[start + 1:]:
            p.paragraph_format.line_spacing = 1.0
    _edit(good, change)
    sub = _sub(check_spacing(good), "double_spacing")
    assert sub["status"] == "fail"
    assert len(sub["problems"]) == 3


def test_tables_are_skipped_for_spacing(good):
    def change(doc):
        cell = doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
        cell.add_run("single spaced cell")
        cell.paragraph_format.line_spacing = 1.0
    _edit(good, change)
    assert check_spacing(good)["status"] == "pass"


# ---------------- APA-03 alignment ----------------

def test_justified_body_paragraph_fails(good):
    _edit(good, lambda d: setattr(_find(d, BODY), "alignment", WD_ALIGN_PARAGRAPH.JUSTIFY))
    sub = _sub(check_alignment(good), "no_justification")
    assert sub["status"] == "fail"
    assert [p["text"][:14] for p in sub["problems"]] == [BODY]


def test_justified_normal_style_fails(good):
    # K-I2 recipe: justification set at style level.
    _edit(good, lambda d: setattr(d.styles["Normal"].paragraph_format, "alignment",
                                  WD_ALIGN_PARAGRAPH.JUSTIFY))
    sub = _sub(check_alignment(good), "no_justification")
    assert sub["status"] == "fail"
    assert len(sub["problems"]) > 5


def test_justified_table_cell_fails(good):
    def change(doc):
        cell = doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
        cell.add_run("justified cell")
        cell.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _edit(good, change)
    sub = _sub(check_alignment(good), "no_justification")
    assert [p["in_table"] for p in sub["problems"]] == [True]


# ---------------- APA-04 indent ----------------

def test_missing_first_line_indent_fails(good):
    _edit(good, lambda d: setattr(d.styles["Normal"].paragraph_format, "first_line_indent",
                                  Inches(0)))
    sub = _sub(check_indent(good), "first_line_half_inch")
    assert sub["status"] == "fail"
    assert sub["problems"][0]["first_line_in"] == 0.0


def test_hanging_body_paragraph_fails(good):
    def change(doc):
        pf = _find(doc, BODY).paragraph_format
        pf.left_indent, pf.first_line_indent = Inches(0.5), -Inches(0.5)
    _edit(good, change)
    problems = _sub(check_indent(good), "first_line_half_inch")["problems"]
    assert [(p["text"][:14], p["first_line_in"]) for p in problems] == [(BODY, -0.5)]


@pytest.mark.parametrize("text, left_in", [
    ("A quoted passage of more than forty words would sit here.", 0.5),   # block quote
    ("Table 1", 0),
    ("Figure 2", 0),
    ("Note. Data are fictional.", 0),
])
def test_non_body_paragraphs_are_not_flagged(good, text, left_in):
    def change(doc):
        para = _new_paragraph(_find(doc, BODY), text)
        para.paragraph_format.first_line_indent = Inches(0)
        para.paragraph_format.left_indent = Inches(left_in)
    _edit(good, change)
    assert check_indent(good)["status"] == "pass"


# ---------------- registry + CLI ----------------

def test_registry_names_and_subset(good):
    assert list(CHECKS) == ["margins", "spacing", "align", "indent"]
    assert set(check_layout(good, ["margins", "align"])) == {"APA-00", "APA-03"}
    with pytest.raises(ValueError):
        check_layout(good, ["nope"])


def test_cli_prints_json(good, capsys):
    assert main(["--docx", str(good), "--checks", "margins,indent"]) == 0
    result = json.loads(capsys.readouterr().out)
    assert {row["status"] for row in result.values()} == {"pass"}
