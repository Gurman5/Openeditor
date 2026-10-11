"""Tests for tools/corpus/checks/preservation.py (K-H9, hard preservation flag).

The "input" is the K-0.3a fixture (sometimes with author noise added); the
"output" is a copy changed the way the engine might change it.
These tests call the real app normalise_docx (read-only import).
Run:  python -m pytest tests/corpus/test_preservation.py -v
"""

import json
import shutil

import pytest
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from tools.corpus.checks.preservation import ENV_VAR, check_preservation, main
from tools.corpus.fixtures import build_good_docx

BODY = "Forty students"          # a plain body paragraph in the fixture


@pytest.fixture
def pair(tmp_path):
    """(input path, output path): two identical copies of the fixture."""
    inp = build_good_docx(tmp_path / "input.docx")
    out = tmp_path / "output.docx"
    shutil.copyfile(inp, out)
    return inp, out


def _edit(path, change):
    doc = Document(path)
    change(doc)
    doc.save(path)
    return path


def _find(doc, start):
    return next(p for p in doc.paragraphs if p.text.startswith(start))


def _wrap_run(run, tag, author="Engine"):
    """Wrap an existing run in a tracked <w:ins> or <w:del> (w:t -> w:delText for del)."""
    wrapper = OxmlElement(tag)
    wrapper.set(qn("w:id"), "901")
    wrapper.set(qn("w:author"), author)
    run._r.addprevious(wrapper)
    wrapper.append(run._r)
    if tag == "w:del":
        for t in run._r.findall(qn("w:t")):
            t.tag = qn("w:delText")


def test_identical_output_passes(pair):
    inp, out = pair
    result = check_preservation(out, inp)
    assert (result["status"], result["hard_flag"], result["ratio"]) == ("pass", False, 0.0)
    assert result["matrix_id"] is None
    assert result["normalisation"]["accepted_revisions"] == 0


def test_small_edit_passes_with_evidence(pair):
    inp, out = pair
    _edit(out, lambda d: setattr(_find(d, BODY).runs[0], "text",
                                 _find(d, BODY).text.replace("voluntary", "optional")))
    result = check_preservation(out, inp)
    assert result["status"] == "pass"
    assert 0 < result["ratio"] < 0.01
    assert (result["inserted_words"], result["deleted_words"]) == (1, 1)
    change = result["biggest_changes"][0]
    assert (change["before"], change["after"]) == ("voluntary", "optional")


def test_losing_half_the_text_raises_hard_flag(pair):
    inp, out = pair
    def change(doc):
        paras = [p for p in doc.paragraphs if p.text.strip()]
        for p in paras[len(paras) // 2:]:
            p._p.getparent().remove(p._p)
    _edit(out, change)
    result = check_preservation(out, inp)
    assert (result["status"], result["hard_flag"]) == ("fail", True)
    assert result["ratio"] > 0.10


def test_tracked_deletion_in_output_counts_as_gone(pair):
    inp, out = pair
    def change(doc):
        p = _find(doc, BODY)
        text = p.runs[0].text
        p.runs[0].text = text.replace(" and all students gave consent.", "")
        tail = p.add_run(" and all students gave consent.")
        _wrap_run(tail, "w:del")
    _edit(out, change)
    result = check_preservation(out, inp)
    assert result["deleted_words"] == 5
    assert result["inserted_words"] == 0


def test_tracked_insertion_in_output_counts_as_present(pair):
    inp, out = pair
    _edit(out, lambda d: _wrap_run(_find(d, BODY).add_run(" Ethics approval was obtained."),
                                   "w:ins"))
    assert check_preservation(out, inp)["inserted_words"] == 4


# ---------------- author noise in the INPUT is not counted ----------------

def test_author_tracked_changes_in_input_are_accepted_first(pair):
    inp, out = pair
    _edit(inp, lambda d: _wrap_run(_find(d, BODY).add_run(" Author added this."), "w:ins",
                                   author="Author"))
    _edit(out, lambda d: _find(d, BODY).add_run(" Author added this."))   # engine accepted it
    result = check_preservation(out, inp)
    assert result["ratio"] == 0.0
    assert result["normalisation"]["accepted_revisions"] >= 1


def test_colour_highlight_and_blank_paragraphs_in_input_are_ignored(pair):
    inp, out = pair
    def change(doc):
        run = _find(doc, BODY).runs[0]
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        _find(doc, "Method").insert_paragraph_before("")      # blank paragraph
    _edit(inp, change)
    result = check_preservation(out, inp)
    assert result["ratio"] == 0.0
    assert result["normalisation"]["runs_highlight_stripped"] >= 1
    assert result["normalisation"]["blank_paragraphs_removed"] >= 1


def test_guidance_notes_paragraph_in_input_is_not_lost_text(pair):
    inp, out = pair
    def change(doc):
        note = _find(doc, BODY).insert_paragraph_before("Delete this guidance before submitting.")
        style = OxmlElement("w:pStyle")
        style.set(qn("w:val"), "GuidanceNotes")
        note._p.get_or_add_pPr().insert(0, style)
    _edit(inp, change)
    result = check_preservation(out, inp)
    assert result["ratio"] == 0.0
    assert result["normalisation"]["guidance_paragraphs_deleted"] == 1


# ---------------- settings + options ----------------

def test_limit_can_be_changed(pair, monkeypatch):
    inp, out = pair
    _edit(out, lambda d: _find(d, BODY)._p.getparent().remove(_find(d, BODY)._p))
    ratio = check_preservation(out, inp)["ratio"]
    assert check_preservation(out, inp, max_ratio=ratio / 2)["hard_flag"] is True
    monkeypatch.setenv(ENV_VAR, str(ratio / 2))
    assert check_preservation(out, inp)["hard_flag"] is True
    monkeypatch.setenv(ENV_VAR, "not-a-number")
    with pytest.raises(ValueError):
        check_preservation(out, inp)


def test_already_normalised_input_option(pair):
    inp, out = pair
    result = check_preservation(out, normalised=inp)
    assert (result["status"], result["normalisation"]) == ("pass", None)
    with pytest.raises(ValueError):
        check_preservation(out)


def test_empty_input_is_na(tmp_path):
    empty = tmp_path / "empty.docx"
    Document().save(empty)
    assert check_preservation(empty, empty)["status"] == "na"


def test_cli_prints_json(pair, capsys):
    inp, out = pair
    assert main(["--docx", str(out), "--input", str(inp)]) == 0
    assert json.loads(capsys.readouterr().out)["hard_flag"] is False
