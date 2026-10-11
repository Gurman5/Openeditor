"""Tests for tools/corpus/checks/fonts.py (K-H3, APA-02 font check).

Each test starts from the K-0.3a synthetic fixture and changes one thing.
Run:  python -m pytest tests/corpus/test_fonts.py -v
"""

import json
import zipfile

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from tools.corpus.checks.fonts import ENV_VAR, check_fonts, main, parse_allowed
from tools.corpus.fixtures import build_good_docx


@pytest.fixture
def good(tmp_path):
    return build_good_docx(tmp_path / "good.docx")


def _edit(path, change):
    """Open the .docx with python-docx, apply change(doc), save in place."""
    doc = Document(path)
    change(doc)
    doc.save(path)
    return path


def _set_normal(doc, name=None, size=None):
    font = doc.styles["Normal"].font
    if name is not None:
        font.name = name
    if size is not None:
        font.size = Pt(size)


def test_fixture_passes(good):
    result = check_fonts(good)
    assert result["status"] == "pass", result["offending"]
    assert result["matrix_id"] == "APA-02"
    assert result["runs_checked"] > 20
    assert result["families_used"] == {"Calibri": result["runs_checked"]}
    assert result["unresolved"] == []


# K-I2 font recipes: (a) disallowed family, (b) allowed family at the wrong size.

def test_disallowed_family_fails(good):
    _edit(good, lambda d: _set_normal(d, name="Comic Sans MS"))
    result = check_fonts(good)
    assert result["status"] == "fail"
    bad = result["offending"][0]
    assert (bad["font"], bad["reason"]) == ("Comic Sans MS", "font not allowed")
    assert bad["font_source"] == "style:Normal"


def test_allowed_family_wrong_size_fails(good):
    _edit(good, lambda d: _set_normal(d, name="Arial", size=12))
    result = check_fonts(good)
    assert result["status"] == "fail"
    bad = result["offending"][0]
    assert bad["reason"] == "Arial needs 11pt, got 12.0"
    assert bad["size_source"] == "style:Normal"


def test_family_match_ignores_case(good):
    _edit(good, lambda d: _set_normal(d, name="calibri"))
    assert check_fonts(good)["status"] == "pass"


# Changing the allowed list without editing the code.

def test_custom_allowed_list(good):
    assert check_fonts(good, allowed={"Georgia": 11})["status"] == "fail"
    assert check_fonts(good, allowed={"Calibri": 11})["status"] == "pass"


def test_env_var_overrides_default(good, monkeypatch):
    monkeypatch.setenv(ENV_VAR, "Georgia:11, Times New Roman:12")
    result = check_fonts(good)
    assert result["allowed"] == {"Georgia": 11.0, "Times New Roman": 12.0}
    assert result["status"] == "fail"
    monkeypatch.setenv(ENV_VAR, "Calibri:11")
    assert check_fonts(good)["status"] == "pass"


@pytest.mark.parametrize("bad", ["Calibri", "Calibri:big", ":11", ""])
def test_bad_allowed_text_raises(bad):
    with pytest.raises(ValueError):
        parse_allowed(bad)


# Tables: family checked, size not.

def _add_table_run(doc, text, name=None, size=None):
    run = doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0].add_run(text)
    if name:
        run.font.name = name
    if size:
        run.font.size = Pt(size)


def test_table_text_size_not_checked(good):
    _edit(good, lambda d: _add_table_run(d, "small table text", size=9))
    assert check_fonts(good)["status"] == "pass"


def test_table_text_family_still_checked(good):
    _edit(good, lambda d: _add_table_run(d, "odd table text", name="Comic Sans MS"))
    result = check_fonts(good)
    assert result["status"] == "fail"
    assert [b["in_table"] for b in result["offending"]] == [True]


# Theme fonts.

def _theme_normal(doc):
    doc.styles["Normal"].element.rPr.find(qn("w:rFonts")).set(qn("w:asciiTheme"), "minorHAnsi")


def test_theme_font_is_resolved_and_graded(good):
    _edit(good, _theme_normal)      # python-docx's theme minor font is Cambria
    result = check_fonts(good)
    assert result["status"] == "fail"
    bad = result["offending"][0]
    assert bad["font"] == "Cambria"
    assert bad["font_source"] == "style:Normal (theme minorHAnsi)"


def test_unresolvable_font_is_reported_not_failed(good):
    _edit(good, _theme_normal)
    with zipfile.ZipFile(good) as z:
        parts = {n: z.read(n) for n in z.namelist() if n != "word/theme/theme1.xml"}
    with zipfile.ZipFile(good, "w") as z:
        for name, data in parts.items():
            z.writestr(name, data)

    result = check_fonts(good)
    assert result["unresolved"]
    assert result["unresolved"][0]["reason"] == "font could not be resolved"
    assert result["status"] == "pass"   # remaining runs (headings) are Calibri 11


def test_no_text_is_na(tmp_path):
    path = tmp_path / "empty.docx"
    Document().save(path)
    assert check_fonts(path)["status"] == "na"


def test_cli_prints_json(good, capsys):
    assert main(["--docx", str(good), "--allowed", "Calibri:11"]) == 0
    result = json.loads(capsys.readouterr().out)
    assert result["status"] == "pass"
