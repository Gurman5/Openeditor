"""Tests for tools/corpus/inject.py + tools/corpus/defects/format_defects.py (K-I1).

The base is always the K-0.3a synthetic fixture.
Run:  python -m pytest tests/corpus/test_inject.py -v
"""

import json

import pytest
from docx import Document

from tools.corpus.checks.abstract import check_abstract
from tools.corpus.checks.fonts import check_fonts
from tools.corpus.checks.headings import check_headings
from tools.corpus.checks.layout import check_layout
from tools.corpus.checks.refs import check_references
from tools.corpus.docx_view import accepted_view
from tools.corpus.fixtures import build_good_docx
from tools.corpus.inject import RECIPES, InjectionError, inject, main


@pytest.fixture
def base(tmp_path):
    return build_good_docx(tmp_path / "good.docx")


def all_statuses(path) -> dict:
    """{matrix_id: status} from every check we have (preservation excluded)."""
    view = accepted_view(path)
    rows = {"APA-02": check_fonts(view)}
    for result in (check_abstract(view), check_headings(view), check_layout(view),
                   check_references(view)):
        rows.update(result)
    return {mid: row["status"] for mid, row in rows.items()}


def _changed(before: dict, after: dict) -> set:
    return {mid for mid in before if before[mid] != after[mid]}


def test_registry_comes_from_format_defects():
    from tools.corpus.defects.format_defects import FORMAT_RECIPES
    assert set(RECIPES) == {"APA-03", "APA-00", "APA-01", "APA-04", "APA-02a", "APA-02b",
                            "APA-05"}
    assert RECIPES["APA-03"] is FORMAT_RECIPES["APA-03"]
    assert {RECIPES[r].matrix_id for r in ("APA-02a", "APA-02b")} == {"APA-02"}


@pytest.mark.parametrize("where", ["style", "paragraph"])
def test_justified_breaks_exactly_apa03(base, tmp_path, where):
    out = tmp_path / f"bad_{where}.docx"
    row = inject(base, ["APA-03"], out, where=where)
    before, after = all_statuses(base), all_statuses(out)
    assert after["APA-03"] == "fail"
    assert _changed(before, after) == {"APA-03"}
    assert row["defects_injected"] == ["APA-03"]
    assert row["verdicts"] == {"APA-03": "fail"}
    assert row["injections"][0]["where"] == where


def test_apa03_defaults_to_paragraph(base, tmp_path):
    row = inject(base, ["APA-03"], tmp_path / "default.docx")
    assert row["where"] is None                          # nothing requested
    assert row["injections"][0]["where"] == "paragraph"  # recipe default used


def test_where_applies_only_to_recipes_that_support_it(base, tmp_path):
    row = inject(base, ["APA-03", "APA-00"], tmp_path / "mixed.docx", where="style")
    assert [i["where"] for i in row["injections"]] == ["style", "section"]
    with pytest.raises(InjectionError, match="supports where='style'"):
        inject(base, ["APA-00"], tmp_path / "x.docx", where="style")


def test_style_justify_keeps_inheriting_heading_aligned(base, tmp_path):
    # Make Heading 2 inherit its (left) alignment from Normal, the case the
    # pinning protects: without it, Heading 2 would become justified too.
    doc = Document(base)
    doc.styles["Heading 2"].paragraph_format.alignment = None
    doc.save(base)

    out = tmp_path / "bad.docx"
    row = inject(base, ["APA-03"], out, where="style")
    assert row["injections"][0]["pinned_styles"]["Heading2"] == "left"
    after = all_statuses(out)
    assert after["APA-11"] == "pass"            # headings unaffected
    assert _changed(all_statuses(base), after) == {"APA-03"}


def test_paragraph_choice_follows_the_seed(base, tmp_path):
    def chosen(seed, name):
        row = inject(base, ["APA-03"], tmp_path / name, seed=seed, where="paragraph")
        return row["injections"][0]["paragraphs"]

    first = chosen(0, "a.docx")
    assert chosen(0, "b.docx") == first
    assert len(first) == 3
    assert any(chosen(s, f"s{s}.docx") != first for s in range(1, 10))


def test_same_inputs_give_byte_identical_output(base, tmp_path):
    a = inject(base, ["APA-03"], tmp_path / "a.docx", where="paragraph", seed=3)
    b = inject(base, ["APA-03"], tmp_path / "b.docx", where="paragraph", seed=3)
    assert a["sha256"] == b["sha256"]
    assert (tmp_path / "a.docx").read_bytes() == (tmp_path / "b.docx").read_bytes()


def test_margin_recipe_breaks_exactly_apa00(base, tmp_path):
    out = tmp_path / "margin.docx"
    row = inject(base, ["APA-00"], out)
    assert _changed(all_statuses(base), all_statuses(out)) == {"APA-00"}
    assert row["injections"][0]["twips"] == "1080"


def test_composite_breaks_both_rows(base, tmp_path):
    out = tmp_path / "composite.docx"
    row = inject(base, ["APA-03", "APA-00"], out)
    assert _changed(all_statuses(base), all_statuses(out)) == {"APA-03", "APA-00"}
    assert row["id"] == "good__APA-03+APA-00__s0"
    assert row["verdicts"] == {"APA-03": "fail", "APA-00": "fail"}
    assert [i["matrix_id"] for i in row["injections"]] == ["APA-03", "APA-00"]


# ---------------- K-I2 formatting recipes ----------------

@pytest.mark.parametrize("recipe, where, row", [
    ("APA-01", "paragraph", "APA-01"),
    ("APA-01", "style", "APA-01"),
    ("APA-04", "paragraph", "APA-04"),
    ("APA-04", "style", "APA-04"),
    ("APA-02a", "paragraph", "APA-02"),
    ("APA-02a", "style", "APA-02"),
    ("APA-02b", "paragraph", "APA-02"),
    ("APA-02b", "style", "APA-02"),
    ("APA-05", None, "APA-05"),
])
def test_k_i2_recipe_breaks_exactly_its_row(base, tmp_path, recipe, where, row):
    out = tmp_path / f"{recipe}_{where}.docx"
    result = inject(base, [recipe], out, where=where)
    assert _changed(all_statuses(base), all_statuses(out)) == {row}
    assert result["defects_injected"] == [row]
    assert result["injections"][0]["recipe_id"] == recipe
    assert result["injections"][0]["where"] == (where or RECIPES[recipe].default_where)


def test_spacing_style_mode_pins_reference_entries(base, tmp_path):
    out = tmp_path / "spacing_style.docx"
    row = inject(base, ["APA-01"], out, where="style")
    assert len(row["injections"][0]["pinned_paragraphs"]) == 3      # the 3 references
    assert all_statuses(out)["APA-12"] == "pass"


def test_indent_style_mode_pins_keywords_and_run_in_headings(base, tmp_path):
    # Add a Level 4 run-in heading so the run-in pinning is exercised too.
    doc = Document(base)
    anchor = next(p for p in doc.paragraphs if p.text == "Participants")
    para = anchor.insert_paragraph_before("")
    anchor._p.addnext(para._p)                     # move it to just after "Participants"
    para.add_run("Recruitment Method.").bold = True
    para.add_run(" Students were invited by email.")
    doc.save(base)
    before = all_statuses(base)
    assert before["APA-11"] == "pass"

    out = tmp_path / "indent_style.docx"
    row = inject(base, ["APA-04"], out, where="style")
    assert len(row["injections"][0]["pinned_paragraphs"]) == 2     # keywords + run-in
    after = all_statuses(out)
    assert (after["APA-09"], after["APA-11"]) == ("pass", "pass")
    assert _changed(before, after) == {"APA-04"}


def test_page_number_recipe_strips_both_field_forms(base, tmp_path):
    row = inject(base, ["APA-05"], tmp_path / "nopage.docx")
    assert row["injections"][0]["fields_removed"] == 2              # fldSimple + fldChar
    assert len(row["injections"][0]["parts"]) == 2


def test_two_recipes_for_one_row_are_refused(base, tmp_path):
    with pytest.raises(InjectionError, match="one defect per matrix row"):
        inject(base, ["APA-02a", "APA-02b"], tmp_path / "x.docx")


def test_k_i2_composite_breaks_exactly_those_rows(base, tmp_path):
    out = tmp_path / "composite_k_i2.docx"
    row = inject(base, ["APA-01", "APA-04", "APA-05"], out)
    assert _changed(all_statuses(base), all_statuses(out)) == {"APA-01", "APA-04", "APA-05"}
    assert row["verdicts"] == {"APA-01": "fail", "APA-04": "fail", "APA-05": "fail"}


def test_header_edit_is_byte_identical(base, tmp_path):
    a = inject(base, ["APA-05", "APA-02b"], tmp_path / "a.docx", seed=4)
    b = inject(base, ["APA-05", "APA-02b"], tmp_path / "b.docx", seed=4)
    assert a["sha256"] == b["sha256"]


# ---------------- refusals ----------------

@pytest.mark.parametrize("defects, message", [
    (["APA-99"], "no recipe"),
    (["APA-03", "APA-03"], "repeated"),
    ([], "at least one"),
])
def test_bad_requests_are_refused(base, tmp_path, defects, message):
    with pytest.raises(InjectionError, match=message):
        inject(base, defects, tmp_path / "x.docx")


def test_base_already_failing_is_refused(base, tmp_path):
    first = tmp_path / "already_bad.docx"
    inject(base, ["APA-03"], first)
    with pytest.raises(InjectionError, match="must pass APA-03"):
        inject(first, ["APA-03"], tmp_path / "again.docx")


@pytest.mark.parametrize("name", ["good.docx", "manifest.json", "MANIFEST.JSON"])
def test_dangerous_out_paths_are_refused(base, tmp_path, name):
    before = (tmp_path / "good.docx").read_bytes()
    with pytest.raises(InjectionError):
        inject(base, ["APA-03"], tmp_path / name)
    assert (tmp_path / "good.docx").read_bytes() == before


def test_only_the_output_file_is_written(base, tmp_path):
    out_dir = tmp_path / "out"
    inject(base, ["APA-03"], out_dir / "bad.docx")
    assert [p.name for p in out_dir.iterdir()] == ["bad.docx"]
    assert not list(tmp_path.rglob("manifest.json"))


def test_cli_prints_the_row(base, tmp_path, capsys):
    out = tmp_path / "cli.docx"
    assert main(["--base", str(base), "--defect", "APA-03", "--out", str(out),
                 "--where", "paragraph", "--seed", "1"]) == 0
    row = json.loads(capsys.readouterr().out)
    assert (row["id"], row["where"]) == ("good__APA-03__s1", "paragraph")
