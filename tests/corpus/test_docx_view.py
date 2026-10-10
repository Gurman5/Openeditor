"""Tests for tools/corpus/docx_view.py (K-0.3 view half + K-H1 validity).

Test documents are built here at test time; nothing is committed.
Run:  python -m pytest tests/corpus/test_docx_view.py -v
"""

import zipfile

import pytest
from docx import Document
from lxml import etree

from tools.corpus.docx_view import accepted_view, check_validity, w

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _make_docx(path, body_xml: str) -> str:
    """Create a real .docx with python-docx, then swap in our own <w:body> XML."""
    doc = Document()
    doc.add_paragraph("placeholder")
    doc.save(path)

    with zipfile.ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist()}

    root = etree.fromstring(parts["word/document.xml"])
    body = root.find(w("body"))
    sect = body.find(w("sectPr"))
    for child in list(body):
        body.remove(child)
    new_body = etree.fromstring(f'<w:body xmlns:w="{W_NS}">{body_xml}</w:body>')
    for child in new_body:
        body.append(child)
    if sect is not None:
        body.append(sect)
    parts["word/document.xml"] = etree.tostring(root, xml_declaration=True,
                                                encoding="UTF-8", standalone=True)

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in parts.items():
            z.writestr(name, data)
    return str(path)


def _r(text):
    return f'<w:r><w:t xml:space="preserve">{text}</w:t></w:r>'


def test_plain_text(tmp_path):
    f = _make_docx(tmp_path / "a.docx", f"<w:p>{_r('Hello world')}</w:p>")
    view = accepted_view(f)
    assert [p.text for p in view.paragraphs] == ["Hello world"]


def test_insertion_kept_deletion_dropped(tmp_path):
    xml = (
        "<w:p>"
        + _r("The cat ")
        + f'<w:del w:id="1" w:author="Bot"><w:r><w:delText>sat</w:delText></w:r></w:del>'
        + f'<w:ins w:id="2" w:author="Bot">{_r("slept")}</w:ins>'
        + _r(" here.")
        + "</w:p>"
    )
    view = accepted_view(_make_docx(tmp_path / "b.docx", xml))
    assert view.paragraphs[0].text == "The cat slept here."
    assert "sat" not in view.text


def test_moves(tmp_path):
    xml = (
        f'<w:p><w:moveFrom w:id="3" w:author="Bot">{_r("old place")}</w:moveFrom>{_r("A")}</w:p>'
        f'<w:p>{_r("B")}<w:moveTo w:id="4" w:author="Bot">{_r(" new place")}</w:moveTo></w:p>'
    )
    view = accepted_view(_make_docx(tmp_path / "c.docx", xml))
    assert [p.text for p in view.paragraphs] == ["A", "B new place"]


def test_formatting_snapshot_ignored(tmp_path):
    # Current spacing is 480; the pPrChange snapshot (old value 276) must be ignored.
    xml = (
        "<w:p><w:pPr><w:spacing w:line=\"480\" w:lineRule=\"auto\"/>"
        "<w:pPrChange w:id=\"5\" w:author=\"Bot\"><w:pPr><w:spacing w:line=\"276\"/></w:pPr></w:pPrChange>"
        "</w:pPr>" + _r("Body") + "</w:p>"
    )
    view = accepted_view(_make_docx(tmp_path / "d.docx", xml))
    ppr = view.paragraphs[0].pPr
    lines = [s.get(w("line")) for s in ppr.iter(w("spacing"))]
    assert lines == ["480"]


def test_deleted_paragraph_mark_merges(tmp_path):
    xml = (
        '<w:p><w:pPr><w:rPr><w:del w:id="6" w:author="Bot"/></w:rPr></w:pPr>'
        + _r("First half, ") + "</w:p>"
        + '<w:p><w:pPr><w:pStyle w:val="BodyText"/></w:pPr>' + _r("second half.") + "</w:p>"
    )
    view = accepted_view(_make_docx(tmp_path / "e.docx", xml))
    assert [p.text for p in view.paragraphs] == ["First half, second half."]
    assert view.paragraphs[0].style_id == "BodyText"


def test_tabs_breaks_and_field_instructions(tmp_path):
    xml = (
        "<w:p><w:r><w:t>A</w:t><w:tab/><w:t>B</w:t><w:br/><w:t>C</w:t></w:r>"
        '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        '<w:r><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>'
    )
    view = accepted_view(_make_docx(tmp_path / "f.docx", xml))
    assert view.paragraphs[0].text == "A\tB\nC"


def test_table_and_style(tmp_path):
    xml = (
        '<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr>' + _r("Method") + "</w:p>"
        "<w:tbl><w:tr><w:tc><w:p>" + _r("cell") + "</w:p></w:tc></w:tr></w:tbl>"
    )
    view = accepted_view(_make_docx(tmp_path / "g.docx", xml))
    assert view.paragraphs[0].style_id == "Heading1"
    assert view.paragraphs[0].in_table is False
    assert view.paragraphs[1].text == "cell"
    assert view.paragraphs[1].in_table is True


def test_original_file_not_changed(tmp_path):
    xml = "<w:p>" + f'<w:del w:id="1" w:author="Bot"><w:r><w:delText>x</w:delText></w:r></w:del>' + "</w:p>"
    f = _make_docx(tmp_path / "h.docx", xml)
    before = (tmp_path / "h.docx").read_bytes()
    accepted_view(f)
    assert (tmp_path / "h.docx").read_bytes() == before


# ---------------- K-H1 validity ----------------

def test_validity_good(tmp_path):
    f = _make_docx(tmp_path / "v.docx", f"<w:p>{_r('ok')}</w:p>")
    result = check_validity(f)
    assert result["valid"] is True
    assert result["duplicate_id_count"] == 0


def test_validity_reports_duplicate_ids(tmp_path):
    xml = (
        f'<w:p><w:ins w:id="7" w:author="Bot">{_r("a")}</w:ins>'
        f'<w:ins w:id="7" w:author="Bot">{_r("b")}</w:ins></w:p>'
    )
    result = check_validity(_make_docx(tmp_path / "dup.docx", xml))
    assert result["valid"] is True          # reported, not failed
    assert result["duplicate_id_count"] == 1
    assert result["duplicate_ids"] == {"7": 2}


def test_validity_not_a_zip(tmp_path):
    bad = tmp_path / "bad.docx"
    bad.write_bytes(b"not a zip")
    result = check_validity(bad)
    assert result["valid"] is False


def test_validity_malformed_xml(tmp_path):
    f = _make_docx(tmp_path / "m.docx", f"<w:p>{_r('ok')}</w:p>")
    with zipfile.ZipFile(f) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    parts["word/document.xml"] = b"<broken"
    with zipfile.ZipFile(f, "w") as z:
        for n, d in parts.items():
            z.writestr(n, d)
    result = check_validity(f)
    assert result["valid"] is False
    assert any("malformed" in r for r in result["reasons"])
