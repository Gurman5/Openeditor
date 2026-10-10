"""Tests for tools/corpus/fixtures.py (K-0.3a synthetic known-good fixture).

The harness (tools/corpus/harness.py) does not exist yet, so these tests read
the raw OOXML directly. When it lands, add one test that runs the core checks
on this fixture and expects all of them to pass.
Run:  python -m pytest tests/corpus/test_fixtures.py -v
"""

import pytest
from docx import Document

from tools.corpus.docx_view import accepted_view, check_validity, parse_xml, read_parts, w
from tools.corpus.fixtures import ABSTRACT, build_good_docx, main

R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


@pytest.fixture
def good_docx(tmp_path):
    return build_good_docx(tmp_path / "synthetic_apa7_good.docx")


def _style(docx_path, style_id):
    styles = parse_xml(read_parts(docx_path)["word/styles.xml"])
    for s in styles.iter(w("style")):
        if s.get(w("styleId")) == style_id:
            return s
    raise AssertionError(f"style {style_id} not found")


def _paragraph(view, text):
    return next(p for p in view.paragraphs if p.text == text)


def _jc(element):
    jc = element.find(".//" + w("jc"))
    return None if jc is None else jc.get(w("val"))


def test_cli_writes_file(tmp_path):
    out = tmp_path / "sub" / "synthetic_apa7_good.docx"
    assert main(["--out", str(out)]) == 0
    assert out.exists()


def test_valid_and_no_tracked_changes(good_docx):
    result = check_validity(good_docx)
    assert result["valid"] is True, result["reasons"]
    assert result["duplicate_id_count"] == 0
    doc_xml = parse_xml(read_parts(good_docx)["word/document.xml"])
    assert not list(doc_xml.iter(w("ins"), w("del"), w("moveFrom"), w("moveTo")))


def test_margins_1440_on_all_sections(good_docx):
    sections = accepted_view(good_docx).sections
    assert sections
    for sect in sections:
        mar = sect.find(w("pgMar"))
        for side in ("top", "bottom", "left", "right"):
            assert mar.get(w(side)) == "1440", side


def test_normal_style_spacing_indent_align(good_docx):
    ppr = _style(good_docx, "Normal").find(w("pPr"))
    spacing = ppr.find(w("spacing"))
    assert spacing.get(w("line")) == "480"
    assert spacing.get(w("lineRule")) == "auto"
    assert ppr.find(w("ind")).get(w("firstLine")) == "720"
    assert _jc(ppr) == "left"


def test_nothing_justified(good_docx):
    parts = read_parts(good_docx)
    for name in ("word/document.xml", "word/styles.xml"):
        root = parse_xml(parts[name])
        assert all(jc.get(w("val")) != "both" for jc in root.iter(w("jc"))), name


@pytest.mark.parametrize("style_id", ["Normal", "Heading1", "Heading2"])
def test_styles_use_calibri_11(good_docx, style_id):
    rpr = _style(good_docx, style_id).find(w("rPr"))
    fonts = rpr.find(w("rFonts"))
    assert fonts.get(w("ascii")) == "Calibri"
    assert fonts.get(w("hAnsi")) == "Calibri"
    assert fonts.get(w("asciiTheme")) is None      # theme font would override ascii
    assert rpr.find(w("sz")).get(w("val")) == "22"  # half-points: 22 = 11 pt
    assert rpr.find(w("color")) is None


def test_no_run_level_font_overrides(good_docx):
    for p in accepted_view(good_docx).paragraphs:
        for r in p.runs:
            if r.rPr is not None:
                assert r.rPr.find(w("rFonts")) is None, p.text
                assert r.rPr.find(w("sz")) is None, p.text


def test_headers_have_page_field_both_forms(good_docx):
    parts = read_parts(good_docx)
    sect = accepted_view(good_docx).sections[-1]
    assert sect.find(w("titlePg")) is not None

    rels = parse_xml(parts["word/_rels/document.xml.rels"])
    targets = {r.get("Id"): r.get("Target") for r in rels.iter(f"{{{PKG_REL_NS}}}Relationship")}
    refs = {ref.get(w("type")): ref.get(f"{{{R_NS}}}id") for ref in sect.iter(w("headerReference"))}
    assert {"default", "first"} <= set(refs)

    first = parse_xml(parts["word/" + targets[refs["first"]]])
    default = parse_xml(parts["word/" + targets[refs["default"]]])
    assert any(f.get(w("instr")).strip() == "PAGE" for f in first.iter(w("fldSimple")))
    assert any(t.text.strip() == "PAGE" for t in default.iter(w("instrText")))
    assert [c.get(w("fldCharType")) for c in default.iter(w("fldChar"))] == ["begin", "separate", "end"]


def test_heading_levels(good_docx):
    view = accepted_view(good_docx)
    h1 = [p.text for p in view.paragraphs if p.style_id == "Heading1"]
    h2 = [p.text for p in view.paragraphs if p.style_id == "Heading2"]
    assert h1 == ["Method", "Results", "Discussion"]
    assert h2 == ["Participants", "Procedure"]
    assert "Introduction" not in h1 + h2

    for style_id, align in (("Heading1", "center"), ("Heading2", "left")):
        style = _style(good_docx, style_id)
        assert _jc(style.find(w("pPr"))) == align
        assert style.find(w("rPr")).find(w("b")) is not None


def test_abstract(good_docx):
    view = accepted_view(good_docx)
    label = _paragraph(view, "Abstract")
    assert label.pPr.find(w("pageBreakBefore")) is not None
    assert _jc(label.pPr) == "center"
    assert label.runs[0].rPr.find(w("b")) is not None

    body = view.paragraphs[label.index + 1]
    assert body.text == ABSTRACT
    assert len(body.text.split()) <= 250
    assert body.pPr.find(w("ind")).get(w("firstLine")) == "0"

    keywords = view.paragraphs[label.index + 2]
    assert keywords.text.startswith("Keywords:")
    assert keywords.runs[0].rPr.find(w("i")) is not None


def test_references(good_docx):
    view = accepted_view(good_docx)
    label = _paragraph(view, "References")
    assert label.pPr.find(w("pageBreakBefore")) is not None
    assert _jc(label.pPr) == "center"
    assert label.runs[0].rPr.find(w("b")) is not None

    entries = view.paragraphs[label.index + 1:]
    assert len(entries) == 3
    for p in entries:
        ind = p.pPr.find(w("ind"))
        assert ind.get(w("left")) == "720"
        assert ind.get(w("hanging")) == "720"
        assert "https://doi.org/" in p.text
    texts = [p.text.lower() for p in entries]
    assert texts == sorted(texts)


def test_under_upload_gate_word_limit(good_docx):
    words = sum(len(p.text.split()) for p in Document(good_docx).paragraphs)
    assert 0 < words <= 15000
