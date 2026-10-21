"""JB-08: the body first-line-indent pass must not indent table captions,
headings, or subheadings (cell paragraphs were already outside the pass).
"""

from docx import Document
from lxml import etree

from app.services.output_generation_samfix import (
    _apply_tracked_body_format_fixes,
    _load_body_paras_raw,
    _table_caption_paragraph_indices,
)

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WQ = f"{{{W}}}"


def _paragraph_texts(body_paras):
    return ["".join(t.text or "" for t in p.iter(f"{WQ}t")) for p in body_paras]


def _has_first_line_indent(para_el):
    ind = para_el.find(f"{WQ}pPr/{WQ}ind")
    return ind is not None and ind.get(f"{WQ}firstLine") == "720"


def test_table_caption_paragraph_indices_detects_paragraph_above_table(tmp_path):
    input_path = tmp_path / "input.docx"

    doc = Document()
    doc.add_paragraph("Descriptive statistics")  # not a caption label
    doc.add_table(rows=1, cols=1)
    doc.add_paragraph("Body text.")
    doc.save(input_path)

    doc_root, body_paras = _load_body_paras_raw(str(input_path))
    heading_idx = _paragraph_texts(body_paras).index("Descriptive statistics")

    skip = _table_caption_paragraph_indices(doc_root, body_paras)
    assert heading_idx in skip


def test_body_indent_skips_table_caption_but_indents_body(tmp_path):
    input_path = tmp_path / "input.docx"
    output_path = tmp_path / "output.docx"

    doc = Document()
    doc.add_paragraph("Table 1. Results of the experiment")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Header A"
    table.cell(0, 1).text = "Header B"
    doc.add_paragraph("This is a normal body paragraph.")
    doc.save(input_path)

    doc_root, body_paras = _load_body_paras_raw(str(input_path))
    texts = _paragraph_texts(body_paras)
    caption_idx = texts.index("Table 1. Results of the experiment")
    body_idx = texts.index("This is a normal body paragraph.")

    skip = _table_caption_paragraph_indices(doc_root, body_paras)
    assert caption_idx in skip
    assert body_idx not in skip

    fixes = [
        {"paragraph_index": caption_idx, "indent": True},
        {"paragraph_index": body_idx, "indent": True},
    ]
    assert _apply_tracked_body_format_fixes(str(input_path), str(output_path), fixes, 5000) is True

    out_root, out_paras = _load_body_paras_raw(str(output_path))
    assert _has_first_line_indent(out_paras[body_idx]) is True
    assert _has_first_line_indent(out_paras[caption_idx]) is False

    # Table cell paragraphs are never touched (they are not body children).
    for tbl in out_root.iter(f"{WQ}tbl"):
        for cell_para in tbl.iter(f"{WQ}p"):
            assert _has_first_line_indent(cell_para) is False


def test_subheading_below_table_is_indented(tmp_path):
    """JB-08 interpretation lock: the plan's "inside ±table-range" wording is
    resolved as "the caption/heading paragraph sitting directly ABOVE a table".
    A Normal paragraph directly BELOW a table is ordinary body prose and
    receives the first-line indent (below-table paragraphs are skipped only
    when they match the caption pattern)."""
    input_path = tmp_path / "input.docx"
    output_path = tmp_path / "output.docx"

    doc = Document()
    doc.add_paragraph("Table set-up summary")  # heading paragraph above table (Normal, no caption label)
    table = doc.add_table(rows=3, cols=3)
    for row in range(3):
        for col in range(3):
            table.cell(row, col).text = f"Cell {row}-{col}"
    doc.add_paragraph("Subheading below the table")
    doc.add_paragraph("Ordinary body paragraph.")
    doc.save(input_path)

    doc_root, body_paras = _load_body_paras_raw(str(input_path))
    texts = _paragraph_texts(body_paras)
    heading_above_idx = texts.index("Table set-up summary")
    subheading_below_idx = texts.index("Subheading below the table")
    body_idx = texts.index("Ordinary body paragraph.")

    skip = _table_caption_paragraph_indices(doc_root, body_paras)
    assert heading_above_idx in skip
    assert subheading_below_idx not in skip
    assert body_idx not in skip

    fixes = [
        {"paragraph_index": idx, "indent": True}
        for idx in (heading_above_idx, subheading_below_idx, body_idx)
    ]
    assert _apply_tracked_body_format_fixes(str(input_path), str(output_path), fixes, 5200) is True

    out_root, out_paras = _load_body_paras_raw(str(output_path))
    assert _has_first_line_indent(out_paras[heading_above_idx]) is False
    assert _has_first_line_indent(out_paras[subheading_below_idx]) is True
    assert _has_first_line_indent(out_paras[body_idx]) is True


def _cell_paragraphs(root):
    return [
        cell_para
        for tbl in root.iter(f"{WQ}tbl")
        for cell_para in tbl.iter(f"{WQ}p")
    ]


def test_spacing_fixes_leave_table_cell_paragraphs_untouched(tmp_path):
    """C2 acceptance: the spacing/format pass iterates direct w:p children of
    w:body only, so cell paragraphs must be skipped entirely — the serialized
    w:spacing / w:ind XML of every cell paragraph must stay byte-identical
    before and after the pass."""
    input_path = tmp_path / "input.docx"
    output_path = tmp_path / "output.docx"

    doc = Document()
    doc.add_paragraph("Table 1. Descriptive statistics")
    doc.add_paragraph("Intro body paragraph.")
    table = doc.add_table(rows=2, cols=2)
    for row in range(2):
        for col in range(2):
            table.cell(row, col).text = f"Cell {row}-{col}"
    doc.add_paragraph("Closing body paragraph.")
    doc.save(input_path)

    doc_root, body_paras = _load_body_paras_raw(str(input_path))
    texts = _paragraph_texts(body_paras)
    body_indices = [
        i
        for i, text in enumerate(texts)
        if text.startswith("Intro body") or text.startswith("Closing body")
        or text.startswith("Table 1.")
    ]
    fixes = [
        {"paragraph_index": idx, "spacing": True, "line_spacing": True, "indent": True}
        for idx in body_indices
    ]

    before_cells = [etree.tostring(cell_para) for cell_para in _cell_paragraphs(doc_root)]

    result = _apply_tracked_body_format_fixes(str(input_path), str(output_path), fixes, 5300)
    assert result is True

    out_root, out_paras = _load_body_paras_raw(str(output_path))
    after_cells = [etree.tostring(cell_para) for cell_para in _cell_paragraphs(out_root)]

    assert before_cells == after_cells
    assert len(before_cells) == 4

    # Sanity: the pass engaged (spacing set on every fixed paragraph, captions
    # included) and its skip set flipped only the indent flag — for the
    # caption-label paragraph AND for the paragraph directly above the table.
    caption_idx = texts.index("Table 1. Descriptive statistics")
    intro_idx = texts.index("Intro body paragraph.")
    closing_idx = texts.index("Closing body paragraph.")
    for idx in (caption_idx, intro_idx, closing_idx):
        spacing = out_paras[idx].find(f"{WQ}pPr/{WQ}spacing")
        assert spacing is not None and spacing.get(f"{WQ}line") == "480"
    assert _has_first_line_indent(out_paras[caption_idx]) is False
    assert _has_first_line_indent(out_paras[intro_idx]) is False
    assert _has_first_line_indent(out_paras[closing_idx]) is True
