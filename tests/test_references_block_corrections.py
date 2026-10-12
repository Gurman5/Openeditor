"""C-25 evidence tests: the References block reaches its APA-12 form.

Covers the three fixes (Heading 1 promotion, tracked bold/centre fallback,
silent page break), their idempotency, and the APA010/011/012 checks against
a document that violates all three.
"""

import shutil
import zipfile

import pytest
from docx import Document
from docx.shared import Pt
from lxml import etree
from unittest.mock import patch

from app.domain import reporting_ownership as ro
from app.services import output_generation_samfix as S
from app.services.apa7_formatting_checks import (
    check_references_hanging_indent,
    check_references_heading_format,
    check_references_page_break,
)
from app.services.apa7_validator import validate as apa_validate
from app.services.docx_format_helpers import (
    has_tracked_page_break_before,
    normalise_heading_text,
    paragraph_text,
    para_style_id,
)
from app.services.references_block_corrections import (
    apply_references_block_corrections,
)

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WQ = f"{{{W}}}"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"

HEADING1_XML = f'''<w:style xmlns:w="{W}" w:type="paragraph" w:styleId="Heading1">
  <w:name w:val="heading 1"/><w:pPr><w:jc w:val="center"/></w:pPr>
  <w:rPr><w:b/></w:rPr>
</w:style>'''


def _rewrap(path, part):
    """Return the parsed XML of one zip part."""
    with zipfile.ZipFile(str(path)) as z:
        return etree.fromstring(z.read(part))


def _add_heading1_style(src):
    """Ensure styles.xml declares a usable paragraph style named 'heading 1'."""
    out = str(src) + ".h1.docx"
    with zipfile.ZipFile(str(src)) as z:
        members = {i.filename: z.read(i.filename) for i in z.infolist()}
    root = etree.fromstring(members["word/styles.xml"])
    names = [
        s.find(f"{WQ}name").get(f"{WQ}val", "").lower()
        for s in root.findall(f"{WQ}style")
        if s.find(f"{WQ}name") is not None
    ]
    if "heading 1" not in names:
        root.append(etree.fromstring(HEADING1_XML.encode()))
    members["word/styles.xml"] = etree.tostring(
        root, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for n, d in members.items():
            zout.writestr(n, d)
    return out


def _drop_heading1_style(src):
    """Remove every 'heading 1' paragraph style (forces the fallback path)."""
    out = str(src) + ".noh1.docx"
    with zipfile.ZipFile(str(src)) as z:
        members = {i.filename: z.read(i.filename) for i in z.infolist()}
    root = etree.fromstring(members["word/styles.xml"])
    for s in list(root.findall(f"{WQ}style")):
        nm = s.find(f"{WQ}name")
        if nm is not None and nm.get(f"{WQ}val", "").strip().lower() == "heading 1":
            root.remove(s)
    members["word/styles.xml"] = etree.tostring(
        root, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for n, d in members.items():
            zout.writestr(n, d)
    return out


def _doc_with_references(tmp_path, heading_style="Normal", bold_heading=False,
                         page_break=False, hanging=True, name="in.docx"):
    """A minimal body ending in a References block, all rules controllable."""
    doc = Document()
    doc.add_paragraph("Some body text.")
    heading = doc.add_paragraph("References", style=heading_style)
    if bold_heading:
        for r in heading.runs:
            r.bold = True
    ref = doc.add_paragraph("Adams, B. (2020). A study. Journal of Things.")
    ref_pr = ref._p.get_or_add_pPr()
    if hanging:
        ind = etree.SubElement(ref_pr, f"{WQ}ind")
        ind.set(f"{WQ}left", "720")
        ind.set(f"{WQ}hanging", "720")
    else:
        first = etree.SubElement(ref_pr, f"{WQ}ind")
        first.set(f"{WQ}firstLine", "720")
    if page_break:
        pb = etree.SubElement(heading._p.get_or_add_pPr(), f"{WQ}pageBreakBefore")
        pb.set(f"{WQ}val", "false")  # present but OFF -> must still be fixed
    path = tmp_path / name
    doc.save(str(path))
    return path


def _refs_heading_el(path):
    root = _rewrap(path, "word/document.xml")
    body = root.find(f"{WQ}body")
    for p in body.iter(f"{WQ}p"):
        if normalise_heading_text(paragraph_text(p)) == "references":
            return p
    return None


def _statuses(results, rule_id):
    for r in results:
        if r["rule_id"] == rule_id:
            return r["status"]
    return "(absent)"


# ── heading promotion ────────────────────────────────────────────────────

def test_bold_normal_heading_promoted_to_heading1(tmp_path):
    src = _add_heading1_style(_doc_with_references(
        tmp_path, heading_style="Normal", bold_heading=True, hanging=True
    ))
    out = tmp_path / "out.docx"

    _, actions = apply_references_block_corrections(src, str(out))

    assert any(a["kind"] == "heading_style" for a in actions)
    assert _statuses(check_references_heading_format(str(out)), "APA011") == "pass"


def test_heading1_style_id_resolved_case_insensitively(tmp_path):
    """A numeric style id ("3") named "heading 1" still counts as Heading 1."""
    doc = Document()
    doc.add_paragraph("Body.")
    doc.add_paragraph("References", style="Normal")
    src = tmp_path / "in.docx"
    doc.save(str(src))

    out = str(src) + ".numeric.docx"
    with zipfile.ZipFile(str(src)) as z:
        members = {i.filename: z.read(i.filename) for i in z.infolist()}
    root = etree.fromstring(members["word/styles.xml"])
    root.append(etree.fromstring(
        f'<w:style xmlns:w="{W}" w:type="paragraph" w:styleId="3">'
        f'<w:name w:val="heading 1"/></w:style>'.encode()
    ))
    members["word/styles.xml"] = etree.tostring(
        root, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for n, d in members.items():
            zout.writestr(n, d)

    dst = tmp_path / "fixed.docx"
    apply_references_block_corrections(out, str(dst))
    assert _statuses(check_references_heading_format(str(dst)), "APA011") == "pass"


# ── tracked fallback (no Heading 1 available) ────────────────────────────

def test_missing_heading1_falls_back_to_tracked_direct_formatting(tmp_path):
    src = _drop_heading1_style(_doc_with_references(
        tmp_path, heading_style="Normal", bold_heading=False, hanging=True
    ))
    out = tmp_path / "out.docx"

    _, actions = apply_references_block_corrections(src, str(out))

    kinds = {a["kind"] for a in actions}
    assert "heading_style" not in kinds
    assert "heading_format" in kinds
    heading = _refs_heading_el(out)
    assert heading is not None
    # bold applied to the run, centre recorded as a tracked pPrChange
    assert heading.find(f"{WQ}r/{WQ}rPr/{WQ}b") is not None
    p_pr = heading.find(f"{WQ}pPr")
    assert p_pr.find(f"{WQ}jc").get(f"{WQ}val") == "center"
    change = p_pr.find(f"{WQ}pPrChange")
    assert change is not None
    assert change.get(f"{WQ}author")
    # w:pPrChange must remain the LAST child of w:pPr — Word rejects it
    # otherwise — so the later silent page break must be inserted BEFORE it.
    assert etree.QName(list(p_pr)[-1]).localname == "pPrChange"
    assert p_pr.find(f"{WQ}pageBreakBefore") is not None
    assert etree.QName(list(p_pr)[-2]).localname == "pageBreakBefore"


def test_fallback_reuses_existing_pprchange(tmp_path):
    """A paragraph may carry only one w:pPrChange; an existing one is reused
    rather than nesting a second (which Word treats as corrupt)."""
    src = _drop_heading1_style(_doc_with_references(
        tmp_path, heading_style="Normal", hanging=True
    ))
    out = str(src) + ".pre.docx"
    with zipfile.ZipFile(str(src)) as z:
        members = {i.filename: z.read(i.filename) for i in z.infolist()}
    root = etree.fromstring(members["word/document.xml"])
    body = root.find(f"{WQ}body")
    for p in body.iter(f"{WQ}p"):
        if normalise_heading_text(paragraph_text(p)) == "references":
            p_pr = p.find(f"{WQ}pPr")
            change = etree.SubElement(p_pr, f"{WQ}pPrChange")
            change.set(f"{WQ}id", "42")
            change.set(f"{WQ}author", "Prior Pass")
    members["word/document.xml"] = etree.tostring(
        root, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for n, d in members.items():
            zout.writestr(n, d)

    dst = tmp_path / "fixed.docx"
    apply_references_block_corrections(out, str(dst))
    heading = _refs_heading_el(dst)
    p_pr = heading.find(f"{WQ}pPr")
    assert len(p_pr.findall(f"{WQ}pPrChange")) == 1
    assert p_pr.find(f"{WQ}pPrChange").get(f"{WQ}id") == "42"


# ── page break ──────────────────────────────────────────────────────────

def test_page_break_inserted_when_absent(tmp_path):
    src = _add_heading1_style(_doc_with_references(tmp_path, hanging=True))
    out = tmp_path / "out.docx"

    _, actions = apply_references_block_corrections(src, str(out))

    assert any(a["kind"] == "page_break" for a in actions)
    heading = _refs_heading_el(out)
    assert has_tracked_page_break_before(heading)
    # silent: the page break is layout plumbing, not a tracked change
    p_pr = heading.find(f"{WQ}pPr")
    assert p_pr.find(f"{WQ}pPrChange") is None
    assert _statuses(check_references_page_break(str(out)), "APA010") == "pass"


def test_page_break_before_val_false_is_still_fixed(tmp_path):
    """A disabled pageBreakBefore must not read as 'already has one'."""
    src = _add_heading1_style(_doc_with_references(
        tmp_path, page_break=True, hanging=True
    ))
    assert _statuses(check_references_page_break(str(src)), "APA010") == "fail"
    out = tmp_path / "out.docx"
    _, actions = apply_references_block_corrections(src, str(out))
    assert any(a["kind"] == "page_break" for a in actions)
    assert _statuses(check_references_page_break(str(out)), "APA010") == "pass"


def test_existing_trailing_break_run_recognised(tmp_path):
    """An author's <w:br w:type="page"/> run counts as a new page — the fixer
    must not add a second break on top."""
    src = _add_heading1_style(_doc_with_references(tmp_path, hanging=True))
    out = str(src) + ".brrun.docx"
    with zipfile.ZipFile(str(src)) as z:
        members = {i.filename: z.read(i.filename) for i in z.infolist()}
    root = etree.fromstring(members["word/document.xml"])
    for p in root.find(f"{WQ}body").iter(f"{WQ}p"):
        if normalise_heading_text(paragraph_text(p)) == "references":
            r = etree.SubElement(p, f"{WQ}r")
            br = etree.SubElement(r, f"{WQ}br")
            br.set(f"{WQ}type", "page")
    members["word/document.xml"] = etree.tostring(
        root, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for n, d in members.items():
            zout.writestr(n, d)

    dst = tmp_path / "fixed.docx"
    _, actions = apply_references_block_corrections(out, str(dst))
    assert not any(a["kind"] == "page_break" for a in actions)


def test_idempotent_second_run(tmp_path):
    src = _add_heading1_style(_doc_with_references(
        tmp_path, heading_style="Normal", bold_heading=True, hanging=True
    ))
    o1 = tmp_path / "o1.docx"
    apply_references_block_corrections(src, str(o1))
    o2 = tmp_path / "o2.docx"
    _, actions = apply_references_block_corrections(str(o1), str(o2))
    assert actions == []


def test_no_references_heading_is_a_noop(tmp_path):
    doc = Document()
    doc.add_paragraph("Body only.")
    src = tmp_path / "in.docx"
    doc.save(str(src))
    out = tmp_path / "out.docx"
    _, actions = apply_references_block_corrections(str(src), str(out))
    assert actions == []
    assert check_references_page_break(str(out)) == []


def test_alias_heading_found(tmp_path):
    """"Bibliography" / "Reference List" are canonical aliases."""
    for alias in ("Reference List", "Bibliography"):
        doc = Document()
        doc.add_paragraph("Body.")
        doc.add_paragraph(alias, style="Normal")
        src = tmp_path / "in.docx"
        doc.save(str(src))
        with zipfile.ZipFile(str(src)) as z:
            names = z.namelist()
        fixed = _add_heading1_style(src)
        out = tmp_path / "out.docx"
        apply_references_block_corrections(fixed, str(out))
        assert _statuses(check_references_page_break(str(out)), "APA010") == "pass", alias
        assert names


# ── APA012 hanging indent ───────────────────────────────────────────────

def test_missing_hanging_indent_flagged(tmp_path):
    src = _doc_with_references(tmp_path, hanging=False)
    res = check_references_hanging_indent(str(src))
    assert _statuses(res, "APA012") == "fail"
    assert "Adams" in res[0]["message"]


def test_hanging_indent_present_passes(tmp_path):
    src = _doc_with_references(tmp_path, hanging=True)
    assert _statuses(check_references_hanging_indent(str(src)), "APA012") == "pass"


def test_style_inherited_hanging_indent_passes(tmp_path):
    """An entry inheriting the indent from its paragraph style is correct."""
    doc = Document()
    doc.add_paragraph("Body.")
    doc.add_paragraph("References", style="Normal")
    ref = doc.add_paragraph("Adams, B. (2020). A study.")
    ref_pr = ref._p.get_or_add_pPr()
    ps = etree.SubElement(ref_pr, f"{WQ}pStyle")
    ps.set(f"{WQ}val", "RefEntry")
    src = tmp_path / "in.docx"
    doc.save(str(src))

    out = str(src) + ".styled.docx"
    with zipfile.ZipFile(str(src)) as z:
        members = {i.filename: z.read(i.filename) for i in z.infolist()}
    root = etree.fromstring(members["word/styles.xml"])
    root.append(etree.fromstring(
        f'<w:style xmlns:w="{W}" w:type="paragraph" w:styleId="RefEntry">'
        f'<w:name w:val="Reference List Entry"/><w:pPr>'
        f'<w:ind w:left="720" w:hanging="720"/></w:pPr></w:style>'.encode()
    ))
    members["word/styles.xml"] = etree.tostring(
        root, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for n, d in members.items():
            zout.writestr(n, d)

    assert _statuses(check_references_hanging_indent(out), "APA012") == "pass"


# ── reporting ownership ─────────────────────────────────────────────────

@pytest.mark.parametrize("rule_id", ["APA010", "APA011", "APA012"])
def test_references_rules_are_sam_fixed_and_listed(rule_id):
    assert rule_id in ro.SAM_FIXED_RULE_IDS
    entry = ro.FAMILY_VERDICTS["APA"]
    assert rule_id in (entry[2] if len(entry) > 2 else [])


# ── end-to-end ──────────────────────────────────────────────────────────

def _run_build(src, out):
    with patch.object(S, "call_llm_json", side_effect=Exception("skip llm")):
        with patch.object(S, "build_body_edit_plan", return_value={"action": "none", "edits": []}):
            with patch.object(
                S, "build_editorial_review_comment_plan",
                return_value={"action": "none", "comments": []},
            ):
                S.build_edited_document(str(src), str(out))


def test_full_build_reaches_apa12(tmp_path):
    """A fixture with a bold-Normal References heading, no page break and a
    first-line-indented entry fails all three rules before the build and
    passes them after."""
    src = tmp_path / "in.docx"
    doc = Document()
    doc.add_paragraph("Body text.")
    heading = doc.add_paragraph("References", style="Normal")
    for r in heading.runs:
        r.bold = True
    ref = doc.add_paragraph("Adams, B. (2020). A study. Journal of Things.")
    ind = etree.SubElement(ref._p.get_or_add_pPr(), f"{WQ}ind")
    ind.set(f"{WQ}firstLine", "720")
    doc.save(str(src))

    before = apa_validate(str(src))["results"]
    assert _statuses(before, "APA010") == "fail"
    assert _statuses(before, "APA011") == "fail"
    assert _statuses(before, "APA012") == "fail"

    out = tmp_path / "out.docx"
    _run_build(src, out)

    after = apa_validate(str(out))["results"]
    assert _statuses(after, "APA010") == "pass"
    assert _statuses(after, "APA011") == "pass"
    assert _statuses(after, "APA012") == "pass"