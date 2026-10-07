"""D-01 Phase 2 evidence tests: one test per validator rule id, proving the
claimed Sam fix inside build_edited_document.

For every KEPT id in SAM_FIXED_RULE_IDS the test:
  1. builds a .docx violating the rule,
  2. runs build_edited_document,
  3. asserts the same validate() rule passes on the output,
  4. names the Sam function responsible for the fix.

The three excluded ids (FP004, FP007, TAB002) get a negative test asserting
they are NOT classified as Sam-fixed — they are comment-only, because no Sam
pass was found that repairs them.
"""
import copy
import shutil

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from unittest.mock import patch

from app.domain.reporting_ownership import SAM_FIXED_RULE_IDS
from app.services.jutlp_validator import validate
from app.services import output_generation_samfix as S

SRC = "tests/jutlp_sample_docx_test_pack/01_valid_identified.docx"


def _status_of(report, rule_id):
    return {r["rule_id"]: r["status"] for r in report["results"]}.get(rule_id, "(absent)")


def _run_full_build(src, out):
    with patch.object(S, "call_llm_json", side_effect=Exception("skip llm")):
        with patch.object(S, "build_body_edit_plan", return_value={"action": "none", "edits": []}):
            with patch.object(S, "build_editorial_review_comment_plan", return_value={"action": "none", "comments": []}):
                S.build_edited_document(str(src), str(out))


def _add_style(doc, name):
    try:
        doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    except ValueError:
        pass


def _strip_paragraphs_by_style(path, stylename):
    d = Document(str(path))
    for para in list(d.paragraphs):
        if para.style.name == stylename:
            para._element.getparent().remove(para._element)
    d.save(str(path))


# ── test bodies ──────────────────────────────────────────────────────────────

def test_fp001_title_plan_restores_title(tmp_path):
    """FP001 fixed by S._apply_title_plan (tracked change)."""
    src = tmp_path / "in.docx"
    shutil.copy(SRC, src)
    _strip_paragraphs_by_style(src, "Article Title")
    assert _status_of(validate(str(src)), "FP001") == "fail"
    out = tmp_path / "out.docx"
    _run_full_build(src, out)
    assert _status_of(validate(str(out)), "FP001") == "pass"


def test_fp002_author_plan_inserts_authors(tmp_path):
    """FP002 fixed by S._apply_author_plan (insert + comment)."""
    src = tmp_path / "in.docx"
    shutil.copy(SRC, src)
    _strip_paragraphs_by_style(src, "Authors")
    assert _status_of(validate(str(src)), "FP002") == "fail"
    out = tmp_path / "out.docx"
    _run_full_build(src, out)
    assert _status_of(validate(str(out)), "FP002") == "pass"


def test_fp003_author_plan_inserts_affiliations(tmp_path):
    """FP003 fixed by S._apply_author_plan (default affiliations insert)."""
    src = tmp_path / "in.docx"
    shutil.copy(SRC, src)
    # 01 fixture has no affiliations styled block, so FP003 already fails.
    assert _status_of(validate(str(src)), "FP003") == "fail"
    out = tmp_path / "out.docx"
    _run_full_build(src, out)
    assert _status_of(validate(str(out)), "FP003") == "pass"


def test_fp005_practitioner_stub(tmp_path):
    """FP005 fixed by S._apply_missing_practitioner_stub (tracked change)."""
    src = tmp_path / "in.docx"
    shutil.copy(SRC, src)
    assert _status_of(validate(str(src)), "FP005") == "fail"
    out = tmp_path / "out.docx"
    _run_full_build(src, out)
    assert _status_of(validate(str(out)), "FP005") == "pass"


def test_fp006_keywords_stub(tmp_path):
    """FP006 fixed by S._apply_missing_keywords_stub (tracked change)."""
    src = tmp_path / "in.docx"
    shutil.copy(SRC, src)
    assert _status_of(validate(str(src)), "FP006") == "fail"
    out = tmp_path / "out.docx"
    _run_full_build(src, out)
    assert _status_of(validate(str(out)), "FP006") == "pass"


def test_fp011_abstract_plan_merges_paragraphs(tmp_path):
    """FP011 fixed by S._apply_abstract_plan (tracked change)."""
    src = tmp_path / "in.docx"
    shutil.copy(SRC, src)
    d = Document(str(src))
    for para in d.paragraphs:
        if para.text.strip() == "Abstract body.":
            dup = copy.deepcopy(para._element)
            para._element.addnext(dup)
            break
    d.save(str(src))
    d = Document(str(src))
    for i, para in enumerate(d.paragraphs):
        if para.text.strip() == "Abstract body.":
            d.paragraphs[i + 1].runs[0].text = "Second abstract paragraph."
            break
    d.save(str(src))
    assert _status_of(validate(str(src)), "FP011") == "fail"
    out = tmp_path / "out.docx"
    _run_full_build(src, out)
    assert _status_of(validate(str(out)), "FP011") == "pass"


def test_sty003_body_style_fixed(tmp_path):
    """STY003 fixed by S._apply_body_and_reference_style_fixes (tracked style change)."""
    src = tmp_path / "in.docx"
    shutil.copy(SRC, src)
    d = Document(str(src))
    _add_style(d, "Subtitle")
    for para in d.paragraphs:
        if para.text.strip() == "Body text.":
            para.style = d.styles["Subtitle"]
    d.save(str(src))
    assert _status_of(validate(str(src)), "STY003") == "warn"
    out = tmp_path / "out.docx"
    _run_full_build(src, out)
    assert _status_of(validate(str(out)), "STY003") == "pass"


def test_sty004_reference_style_fixed(tmp_path):
    """STY004 fixed by S._apply_body_and_reference_style_fixes (tracked style change)."""
    src = tmp_path / "in.docx"
    shutil.copy(SRC, src)
    d = Document(str(src))
    d.add_paragraph("References", style="Heading 1")
    d.add_paragraph("Smith, J. (2020). Title. Journal.", style="Normal")
    d.save(str(src))
    assert _status_of(validate(str(src)), "STY004") == "warn"
    out = tmp_path / "out.docx"
    _run_full_build(src, out)
    assert _status_of(validate(str(out)), "STY004") == "pass"


def test_spe002_intro_page_break(tmp_path):
    """SPE002 fixed by S._apply_intro_page_break (silent pageBreakBefore)."""
    src = tmp_path / "in.docx"
    shutil.copy(SRC, src)
    assert _status_of(validate(str(src)), "SPE002") == "warn"
    out = tmp_path / "out.docx"
    _run_full_build(src, out)
    assert _status_of(validate(str(out)), "SPE002") == "pass"


def test_tab001_table_title_style(tmp_path):
    """TAB001 fixed by S._apply_tracked_table_formatting (silent style retag)."""
    src = tmp_path / "in.docx"
    shutil.copy(SRC, src)
    d = Document(str(src))
    _add_style(d, "Table Number")
    _add_style(d, "Table Title")
    d.add_paragraph("Table 2", style="Table Number")
    d.add_paragraph("Not a title", style="Normal")
    d.add_table(rows=1, cols=1)
    d.save(str(src))
    assert _status_of(validate(str(src)), "TAB001") == "fail"
    out = tmp_path / "out.docx"
    _run_full_build(src, out)
    assert _status_of(validate(str(out)), "TAB001") == "pass"


# ── negative tests: not Sam-fixed, so NOT in the ownership set ───────────────

@pytest.mark.parametrize("rule_id", ["FP004", "FP007", "TAB002"])
def test_not_sam_fixed(rule_id):
    """FP004, FP007, TAB002 have no demonstrated Sam fix — they must stay
    comment-only and out of SAM_FIXED_RULE_IDS."""
    assert rule_id not in SAM_FIXED_RULE_IDS


def test_sam_fixed_set_matches_evidence():
    assert SAM_FIXED_RULE_IDS == {
        "FP001", "FP002", "FP003", "FP005", "FP006", "FP011",
        "STY003", "STY004", "SPE002", "TAB001",
    }
