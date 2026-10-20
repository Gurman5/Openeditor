"""C-28 / APA-10: a literal 'Introduction' Heading 1 becomes the repeated title."""

import zipfile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from lxml import etree

from app.services.output_generation_samfix import (
    _apply_apa_heading_styles,
    _apply_heading_1_text_normalization,
)
from tools.corpus.checks.headings import check_headings

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WQ = f"{{{W}}}"
NS = {"w": W}


def _make_doc(path, title="a study of AI", body_heading="Introduction"):
    doc = Document()
    for name in ("Article Title", "Authors", "Author Affiliations"):
        doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph(title, style="Article Title")
    doc.add_paragraph("Jane Doe", style="Authors")
    doc.add_paragraph("University of Example", style="Author Affiliations")
    doc.add_paragraph(body_heading, style="Heading 1")
    doc.add_paragraph("Body text.")
    doc.save(path)


def _body_paragraphs(path):
    with zipfile.ZipFile(path) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    return root.xpath("//w:body/w:p", namespaces=NS)


def _paragraph_text(paragraph):
    return "".join(node.text or "" for node in paragraph.iter(f"{WQ}t"))


def _heading_1_texts(path):
    return [
        _paragraph_text(p)
        for p in _body_paragraphs(path)
        if p.find(f"{WQ}pPr/{WQ}pStyle") is not None
        and p.find(f"{WQ}pPr/{WQ}pStyle").get(f"{WQ}val").lower() == "heading1"
    ]


def test_introduction_heading_becomes_repeated_title_and_passes_apa10(tmp_path):
    source = tmp_path / "input.docx"
    styled = tmp_path / "styled.docx"
    output = tmp_path / "output.docx"
    _make_doc(source)

    assert _apply_apa_heading_styles(str(source), str(styled)) is True
    assert _apply_heading_1_text_normalization(str(styled), str(output)) is True

    assert _heading_1_texts(output) == ["a study of AI"]
    assert all(_paragraph_text(p) != "Introduction" for p in _body_paragraphs(output))

    result = check_headings(str(output))
    assert result["APA-10"]["status"] == "pass", result["APA-10"]
    assert result["APA-10"]["subchecks"]["title_repeated"]["bold_centered_copies"]


def test_document_without_introduction_heading_is_left_untouched(tmp_path):
    source = tmp_path / "input.docx"
    styled = tmp_path / "styled.docx"
    output = tmp_path / "output.docx"
    _make_doc(source, body_heading="Method")

    assert _apply_apa_heading_styles(str(source), str(styled)) is True
    assert _apply_heading_1_text_normalization(str(styled), str(output)) is False
    assert output.exists() is False


def test_missing_front_page_title_returns_false(tmp_path):
    source = tmp_path / "input.docx"
    output = tmp_path / "output.docx"

    doc = Document()
    doc.add_paragraph("")
    doc.add_paragraph("Introduction", style="Heading 1")
    doc.save(source)

    assert _apply_heading_1_text_normalization(str(source), str(output)) is False
    assert output.exists() is False
