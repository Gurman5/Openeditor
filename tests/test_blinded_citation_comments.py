"""Tests for the blinded / placeholder citation comment pass."""

import zipfile

from docx import Document
from lxml import etree

from app.services.blinded_citation_comments import apply_blinded_citation_comments

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WQ = f"{{{W}}}"


def _comment_texts(docx_path: str) -> list[str]:
    with zipfile.ZipFile(docx_path, "r") as z:
        if "word/comments.xml" not in z.namelist():
            return []
        root = etree.fromstring(z.read("word/comments.xml"))
    return [
        "".join(t.text or "" for t in c.iter(f"{WQ}t"))
        for c in root.findall(f"{WQ}comment")
    ]


def _doc(tmp_path, paragraphs):
    path = tmp_path / "doc.docx"
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    doc.save(str(path))
    return str(path)


def test_author_under_review_citation_is_flagged(tmp_path):
    src = _doc(tmp_path, [
        "Drawing on prior work ([Author] et al., under review), responses "
        "were coded for attribution type.",
    ])
    out = tmp_path / "out.docx"

    next_id, actions = apply_blinded_citation_comments(src, str(out), 7)

    assert next_id == 7  # comment-only: change id untouched
    assert len(actions) == 1
    texts = _comment_texts(str(out))
    assert len(texts) == 1
    assert "[Author]" in texts[0]
    assert "under review" in texts[0].lower()
    assert "jutlp" in texts[0].lower()


def test_university_placeholder_is_flagged_with_institution_wording(tmp_path):
    src = _doc(tmp_path, [
        "The course was delivered at [University] across three years.",
    ])
    out = tmp_path / "out.docx"

    _, actions = apply_blinded_citation_comments(src, str(out))

    assert len(actions) == 1
    text = _comment_texts(str(out))[0]
    assert "[University]" in text
    assert "institution" in text.lower()


def test_anonymised_prose_is_not_flagged(tmp_path):
    """Unbracketed 'anonymous'/'anonymised' prose must NOT trigger."""
    src = _doc(tmp_path, [
        "Students completed an anonymous feedback survey each semester.",
        "All student data were anonymised before analysis.",
    ])
    out = tmp_path / "out.docx"

    _, actions = apply_blinded_citation_comments(src, str(out))

    assert actions == []
    assert _comment_texts(str(out)) == []


def test_one_comment_per_paragraph_listing_all_placeholders(tmp_path):
    src = _doc(tmp_path, [
        "Implemented at [University], China; design described in "
        "([Author] et al., under review).",
    ])
    out = tmp_path / "out.docx"

    _, actions = apply_blinded_citation_comments(src, str(out))

    assert len(actions) == 1  # one comment, even with two placeholders
    text = _comment_texts(str(out))[0]
    assert "[University]" in text and "[Author]" in text
    # both follow-up instructions present
    assert "self-citation" in text.lower()
    assert "institution" in text.lower()


def test_clean_document_writes_no_comments(tmp_path):
    src = _doc(tmp_path, [
        "Smith and Jones (2024) report a similar effect in their study.",
    ])
    out = tmp_path / "out.docx"

    _, actions = apply_blinded_citation_comments(src, str(out))

    assert actions == []
    assert _comment_texts(str(out)) == []
