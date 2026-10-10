"""Synthetic known-good APA 7 fixture (K-0.3a).

Builds synthetic_apa7_good.docx at test time with python-docx. The builder IS
the fixture: never commit the generated .docx.

What the document has (all formatting lives on styles, like a real paper):
  - Normal style: Calibri 11, double spacing (w:line="480" lineRule auto),
    0.5" first-line indent (720 twips), left aligned, no space before/after
  - 1" margins (1440 twips) on every side
  - "different first page" header (w:titlePg). The first-page header uses a
    w:fldSimple PAGE field and the default header uses the w:fldChar/w:instrText
    form, so both K-H2 detection paths have a real example
  - title page, Abstract (bold centered label, one unindented paragraph),
    Keywords line, body with Level 1 + Level 2 headings, References on a new
    page with 0.5" hanging indents in alphabetical order
  - no tracked changes

Usage:
  python -m tools.corpus.fixtures --out <tmp>/synthetic_apa7_good.docx
"""

from __future__ import annotations

import argparse
import datetime as dt
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

FONT_NAME = "Calibri"
FONT_SIZE_PT = 11

TITLE = "Effects of Spaced Practice on Vocabulary Retention in First-Year Students"

ABSTRACT = (
    "This study examined whether spacing vocabulary practice across several "
    "short sessions improves retention compared with a single massed session. "
    "Forty first-year students were randomly assigned to a spaced or a massed "
    "condition and completed a delayed recall test one week later. Students in "
    "the spaced condition recalled more words than students in the massed "
    "condition, and the difference remained after controlling for prior "
    "vocabulary knowledge. These results suggest that simple changes to study "
    "timing can produce meaningful gains in retention for new university "
    "students."
)

KEYWORDS = "spaced practice, vocabulary, retention, first-year students"

# (heading level or None for body text, text)
BODY = [
    (None, "Many first-year students study new vocabulary in a single long "
           "session before an assessment. Research on memory suggests that "
           "spreading practice across time may lead to better long-term "
           "retention (Adams & Brown, 2019)."),
    (None, "The present study tested this idea with a short classroom "
           "intervention that could be adopted without extra resources "
           "(Chen et al., 2021)."),
    (1, "Method"),
    (2, "Participants"),
    (None, "Forty students enrolled in an introductory course took part. "
           "Participation was voluntary and all students gave consent."),
    (2, "Procedure"),
    (None, "Students in the spaced condition practised for 10 minutes on "
           "three separate days. Students in the massed condition practised "
           "for 30 minutes on one day."),
    (1, "Results"),
    (None, "Students in the spaced condition recalled more words on the "
           "delayed test than students in the massed condition."),
    (1, "Discussion"),
    (None, "Spacing practice improved retention in this sample, consistent "
           "with earlier findings (Davis, 2020)."),
]

# Fictional references, already in alphabetical order:
# (text before the italic part, italic part, text after).
REFERENCES = [
    ("Adams, J., & Brown, K. (2019). Spacing effects in classroom learning. ",
     "Journal of Educational Practice, 12",
     "(3), 45–60. https://doi.org/10.0000/jep.2019.003"),
    ("Chen, L., Ortiz, M., & Patel, R. (2021). Low-cost study interventions for "
     "new university students. ",
     "Higher Education Review, 8",
     "(1), 1–15. https://doi.org/10.0000/her.2021.001"),
    ("Davis, S. (2020). ",
     "Memory and learning in everyday settings",
     ". Example Academic Press. https://doi.org/10.0000/eap.2020"),
]


# --------------------------------------------------------------------------
# Style setup
# --------------------------------------------------------------------------

def _set_style_font(style, bold: bool = False, italic: bool = False) -> None:
    """Give a style an explicit Calibri 11 font, removing the template's theme
    fonts (w:asciiTheme etc.), colour and size so nothing else wins."""
    rpr = style.element.get_or_add_rPr()
    for tag in ("w:rFonts", "w:color", "w:sz", "w:szCs"):
        for el in rpr.findall(qn(tag)):
            rpr.remove(el)
    rfonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT_NAME)
    rpr.insert(0, rfonts)
    style.font.size = Pt(FONT_SIZE_PT)
    style.font.bold = bold
    style.font.italic = italic


def _set_double_spacing(pf) -> None:
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE  # w:line="480" w:lineRule="auto"
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)


def _setup_styles(doc) -> None:
    normal = doc.styles["Normal"]
    _set_style_font(normal)
    pf = normal.paragraph_format
    _set_double_spacing(pf)
    pf.first_line_indent = Inches(0.5)          # 720 twips
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # APA Level 1: centered, bold. Level 2: flush left, bold.
    for style_name, align in (("Heading 1", WD_ALIGN_PARAGRAPH.CENTER),
                              ("Heading 2", WD_ALIGN_PARAGRAPH.LEFT)):
        style = doc.styles[style_name]
        _set_style_font(style, bold=True)
        hpf = style.paragraph_format
        _set_double_spacing(hpf)
        hpf.first_line_indent = Inches(0)
        hpf.alignment = align


# --------------------------------------------------------------------------
# Page setup + headers
# --------------------------------------------------------------------------

def _run(text: str | None = None) -> OxmlElement:
    r = OxmlElement("w:r")
    if text is not None:
        t = OxmlElement("w:t")
        t.text = text
        r.append(t)
    return r


def _add_page_field_simple(paragraph) -> None:
    """PAGE field, short form: <w:fldSimple w:instr=" PAGE ">."""
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), " PAGE ")
    fld.append(_run("1"))
    paragraph._p.append(fld)


def _add_page_field_complex(paragraph) -> None:
    """PAGE field, long form: fldChar begin / instrText / separate / result / end."""
    def fld_char(kind: str):
        r = _run()
        fc = OxmlElement("w:fldChar")
        fc.set(qn("w:fldCharType"), kind)
        r.append(fc)
        return r

    instr_run = _run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    instr_run.append(instr)

    for r in (fld_char("begin"), instr_run, fld_char("separate"),
              _run("1"), fld_char("end")):
        paragraph._p.append(r)


def _setup_page(doc) -> None:
    section = doc.sections[0]
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, side, Inches(1))       # 1440 twips
    section.different_first_page_header_footer = True   # w:titlePg

    for header, add_field in ((section.first_page_header, _add_page_field_simple),
                              (section.header, _add_page_field_complex)):
        p = header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.first_line_indent = Inches(0)
        add_field(p)


# --------------------------------------------------------------------------
# Content
# --------------------------------------------------------------------------

def _centered(doc, text: str, bold: bool = False, page_break_before: bool = False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.page_break_before = page_break_before
    p.add_run(text).bold = bold
    return p


def _add_title_page(doc) -> None:
    _centered(doc, TITLE, bold=True)
    for line in ("Alex Morgan", "School of Education, Example University",
                 "EDU1001: Learning and Memory", "Dr. Sam Lee", "March 4, 2026"):
        _centered(doc, line)


def _add_abstract(doc) -> None:
    _centered(doc, "Abstract", bold=True, page_break_before=True)
    p = doc.add_paragraph(ABSTRACT)
    p.paragraph_format.first_line_indent = Inches(0)   # APA: abstract is not indented

    kw = doc.add_paragraph()                           # keeps Normal's 0.5" indent
    kw.add_run("Keywords:").italic = True
    kw.add_run(" " + KEYWORDS)


def _add_body(doc) -> None:
    _centered(doc, TITLE, bold=True, page_break_before=True)  # title repeated, no "Introduction"
    for level, text in BODY:
        if level is None:
            doc.add_paragraph(text)
        else:
            doc.add_paragraph(text, style=f"Heading {level}")


def _add_references(doc) -> None:
    _centered(doc, "References", bold=True, page_break_before=True)
    for before, italic, after in REFERENCES:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Inches(0.5)
        pf.first_line_indent = -Inches(0.5)     # w:hanging="720"
        p.add_run(before)
        p.add_run(italic).italic = True
        p.add_run(after)


# --------------------------------------------------------------------------
# Public entry points
# --------------------------------------------------------------------------

def build_good_docx(path: str | Path) -> Path:
    """Write the synthetic known-good APA 7 document to `path` and return it."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    _setup_styles(doc)
    _setup_page(doc)
    _add_title_page(doc)
    _add_abstract(doc)
    _add_body(doc)
    _add_references(doc)

    props = doc.core_properties
    fixed = dt.datetime(2026, 1, 1, tzinfo=dt.timezone.utc)
    props.author = props.last_modified_by = "OpenEditor corpus fixture"
    props.title = TITLE
    props.created = props.modified = fixed
    props.revision = 1

    doc.save(path)
    return path


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Build the synthetic known-good APA 7 .docx")
    ap.add_argument("--out", required=True, help="where to write the .docx")
    args = ap.parse_args(argv)
    print(build_good_docx(args.out))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
