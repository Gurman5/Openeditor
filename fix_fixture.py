import os
import docx
from docx.enum.style import WD_STYLE_TYPE

target_folder = r"tests/jutlp_sample_docx_test_pack"
os.makedirs(target_folder, exist_ok=True)

def get_or_add_style(doc, name, parent_name=None):
    if name in doc.styles:
        return doc.styles[name]
    style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    if parent_name and parent_name in doc.styles:
        style.base_style = doc.styles[parent_name]
    return style

def create_base_doc(has_authors=True, deidentified=False):
    doc = docx.Document()

    get_or_add_style(doc, "Article Title")
    get_or_add_style(doc, "Authors")
    get_or_add_style(doc, "Author Affiliations")
    get_or_add_style(doc, "Practitioner Notes")
    get_or_add_style(doc, "List Bullet")

    # Front Matter
    p_title = doc.add_paragraph("A Comprehensive Study on Academic Formatting")
    p_title.style = "Article Title"

    if has_authors:
        p_author = doc.add_paragraph("Jane Doe, John Smith")
        p_author.style = "Authors"

        p_affil = doc.add_paragraph("Department of Education, University of Research")
        p_affil.style = "Author Affiliations"
    elif deidentified:
        p_author = doc.add_paragraph("[De-identified for peer review]")
        p_author.style = "Authors"

        p_affil = doc.add_paragraph("[De-identified for peer review]")
        p_affil.style = "Author Affiliations"

    doc.add_heading("Abstract", level=1)
    doc.add_paragraph("This paper presents an in-depth analysis of document validation rules.")

    doc.add_heading("Practitioner Notes", level=1)
    for i in range(1, 6):
        p = doc.add_paragraph(f"• Key insight for practitioners note number {i}.")
        try:
            p.style = "Practitioner Notes"
        except Exception:
            p.style = "List Bullet"

    doc.add_heading("Keywords", level=1)
    doc.add_paragraph("validation, jutlp, document structure, python, unit testing")

    sections = [
        "Introduction",
        "Literature Review",
        "Methodology",
        "Results",
        "Discussion",
        "Conclusion",
        "Acknowledgements",
        "References",
    ]

    for sec in sections:
        doc.add_heading(sec, level=1)
        doc.add_paragraph(f"This is the body text content for the {sec} section.")

        if sec in ("Methodology", "Method"):
            doc.add_heading("Research Design", level=2)
            doc.add_paragraph("Detailed description of research design.")
            doc.add_heading("Participants", level=2)
            doc.add_paragraph("Detailed description of participants involved in the study.")
            doc.add_heading("Measures", level=2)
            doc.add_paragraph("Detailed description of instruments and materials used.")
            doc.add_heading("Procedure", level=2)
            doc.add_paragraph("Step-by-step description of the procedure followed.")
            doc.add_heading("Analysis", level=2)
            doc.add_paragraph("Explanation of the statistical methods used to analyze data.")

        if sec == "Discussion":
            doc.add_heading("Practical Implications", level=2)
            doc.add_paragraph("Discussion of practical implications.")
            doc.add_heading("Theoretical Implications", level=2)
            doc.add_paragraph("Discussion of theoretical implications.")
            doc.add_heading("Limitations and Future Research", level=2)
            doc.add_paragraph("Overview of limitations and future research directions.")

    return doc

# 1. 01_valid_identified.docx
doc1 = create_base_doc(has_authors=True)
doc1.save(os.path.join(target_folder, "01_valid_identified.docx"))

# 2. 02_missing_method_subsection.docx
doc2 = create_base_doc(has_authors=True)
for p in list(doc2.paragraphs):
    if p.text == "Participants":
        p._element.getparent().remove(p._element)
doc2.save(os.path.join(target_folder, "02_missing_method_subsection.docx"))

# 3. 03_front_page_issues.docx
doc3 = create_base_doc(has_authors=True)
notes_removed = 0
for p in list(doc3.paragraphs):
    if "Key insight for practitioners note" in p.text:
        p._element.getparent().remove(p._element)
        notes_removed += 1
        if notes_removed == 1:
            break

for p in doc3.paragraphs:
    if "validation, jutlp" in p.text:
        p.text = "validation, jutlp, document structure, python, unit testing, extrakeyword"
        break
doc3.save(os.path.join(target_folder, "03_front_page_issues.docx"))

# 4. 04_structure_and_endmatter_issues.docx
doc4 = create_base_doc(has_authors=True)
for p in list(doc4.paragraphs):
    if p.text == "Results":
        p.text = "Results and Discussion"
    elif p.text == "Discussion":
        p._element.getparent().remove(p._element)

get_or_add_style(doc4, "Guidance Notes")
p_gn = doc4.add_paragraph("Guidance note paragraph text.")
p_gn.style = "Guidance Notes"
doc4.save(os.path.join(target_folder, "04_structure_and_endmatter_issues.docx"))

# 5. 05_valid_deidentified.docx
doc5 = create_base_doc(has_authors=False, deidentified=True)
doc5.save(os.path.join(target_folder, "05_valid_deidentified.docx"))

# 6. 06_deidentified_author_leak.docx
doc6 = create_base_doc(has_authors=False, deidentified=True)
for p in doc6.paragraphs:
    if "This is the body text content for the Acknowledgements section" in p.text:
        p.text += " The authors would like to thank Jane Doe for assistance with data collection."
        break
else:
    doc6.add_paragraph("The authors would like to thank Jane Doe for assistance with data collection.")

doc6.save(os.path.join(target_folder, "06_deidentified_author_leak.docx"))

print("Updated test fixtures successfully.")