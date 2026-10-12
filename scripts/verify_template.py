"""Verify the generated APA7 Base Template.docx (C-47 acceptance).

Checks: required styles present with correct Calibri 11 formatting, single
1in pgMar, Normal geometry, PAGE header, non-empty skeleton, and that every
expected paragraph style is actually used by >=1 paragraph. The three
title-page roles are style donors for the C-24 title-page pass.
Note:
this script is not a full C-47 test, but a sanity check for the generated 
template and its future use in the C-27 pipeline.
"""

# TODO: Delete script and repurpose into test after completion of Track C
import os
import sys
import zipfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document
from docx.enum.style import WD_STYLE_TYPE

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
d = Document("app/domain/APA7 Base Template.docx")

ok = True

print("--- required styles ---")
for name in ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3",
             "Heading 4", "Heading 5", "Reference",
             "APA 7 Reference List Entry", "Quote", "Table Number",
             "Table Title", "Table Text", "Table Emphasis", "Table Note",
             "Table Grid", "Figure Number", "Figure Title",
             "Figure/Table Number", "Figure/Table Title", "Figure/Table Notes",
             "Article Title", "Authors", "Author Affiliations"]:
    present = name in d.styles
    print(f"  {name:28} present={present}")
    ok = ok and present
    if present:
        st = d.styles[name]
        if st.type == WD_STYLE_TYPE.PARAGRAPH:
            f = st.font
            size = f.size.pt if f.size else None
            cal = f.name == "Calibri" and size == 11
            print(f"    -> Calibri11={cal} bold={f.bold} italic={f.italic}")
            ok = ok and cal

print("--- char companions (presence) ---")
for name in ["Table Number Char", "Table Title Char", "Figure Number Char",
             "Figure Title Char", "Figure/Table Number Char",
             "Figure/Table Title Char", "Figure/Table Notes Char",
             "Table Emphasis Char", "Table Note Char", "Table Text Char"]:
    present = name in d.styles
    print(f"  {name:28} present={present}")
    ok = ok and present

print("--- APA title-page styles ---")
for name in ["Article Title", "Authors", "Author Affiliations"]:
    present = name in d.styles
    print(f"  {name:28} present={present}")
    ok = ok and present

sect = d.sections[0]
pgmars = sect._sectPr.findall(f".//{{{W}}}pgMar")
print(f"--- pgMar count={len(pgmars)} (must be 1) ---")
ok = ok and len(pgmars) == 1
if pgmars:
    pm = pgmars[0]
    vals = {a: pm.get(f"{{{W}}}{a}") for a in ["top", "right", "bottom", "left"]}
    print("  margins:", vals)
    ok = ok and all(v == "1440" for v in vals.values())

hdr = sect.header.paragraphs[0]
xml = hdr._p.xml
print("--- header ---")
print("  has PAGE field:", "PAGE" in xml)
print("  right aligned:", str(hdr.paragraph_format.alignment))
ok = ok and "PAGE" in xml

first_hdr_xml = sect.first_page_header._element.xml
print("  first-page banner:", "graphic" in first_hdr_xml or "blip" in first_hdr_xml)
ok = ok and ("graphic" in first_hdr_xml or "blip" in first_hdr_xml)

with zipfile.ZipFile("app/domain/APA7 Base Template.docx", "r") as archive:
    no_effects_part = "word/stylesWithEffects.xml" not in archive.namelist()
    print("  stylesWithEffects removed:", no_effects_part)
    ok = ok and no_effects_part

print("--- skeleton ---")
paras = [p for p in d.paragraphs if p.text.strip()]
print(f"  non-empty paragraphs: {len(paras)} (must be > 0)")
ok = ok and len(paras) > 0
used = {p.style.name for p in d.paragraphs}
for name in ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3",
             "Heading 4", "Heading 5", "Quote", "Table Number",
             "Table Title", "Table Note", "Figure Number", "Figure Title",
             "Figure/Table Notes", "APA 7 Reference List Entry",
             "Article Title", "Authors", "Author Affiliations"]:
    is_used = name in used
    print(f"  style used: {name:28} {is_used}")
    ok = ok and is_used
print(f"  tables: {len(d.tables)} (must be >= 1)")
ok = ok and len(d.tables) >= 1

print("RESULT:", "PASS" if ok else "FAIL")
raise SystemExit(0 if ok else 1)
